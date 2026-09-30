"""Read-only Excel snapshot of the active research cost ledger."""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import UTC, datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins


PRICE_BASIS = {
    "latest": "最新价格", "historical_latest": "当期最新价格",
    "current_latest_fallback": "当前最新价格补位",
    "current_inventory_fallback": "当前库存价格补位",
    "inventory": "库存价格", "recipe": "引用产品成本",
    "composite": "复配计算成本", "missing": "缺少价格",
}
STATUS = {
    "ready": "可核算", "missing": "待补价格",
    "updating": "更新中（沿用上次完成记录）",
    "failed": "更新失败（沿用上次完成记录）",
}
NUMBER_FORMAT = '0.00'
CHINA_TIME = timezone(timedelta(hours=8))
FILL_DARK = PatternFill("solid", fgColor="233842")
FILL_LIGHT = PatternFill("solid", fgColor="F2F6F7")
FILL_ALT = PatternFill("solid", fgColor="F8FAFA")
FILL_LATEST = PatternFill("solid", fgColor="F0F6F9")
FILL_INVENTORY = PatternFill("solid", fgColor="F0F8F3")
FILL_GROUP = PatternFill("solid", fgColor="45616B")
FILL_LATEST_GROUP = PatternFill("solid", fgColor="31566A")
FILL_INVENTORY_GROUP = PatternFill("solid", fgColor="316B57")
ROW_BORDER = Border(bottom=Side(style="hair", color="DDE5E8"))


def _text(cell, value):
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"  # A product code or owner must never become an Excel formula.


def _number(cell, value, fmt=NUMBER_FORMAT):
    if value is not None:
        cell.value = float(Decimal(str(value)))
        cell.number_format = fmt


def _cost(cell, value):
    if value is not None:
        _number(cell, Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _recorded_at(cell, value):
    if not value:
        return
    timestamp = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if timestamp.tzinfo is None:
        timestamp = timestamp.replace(tzinfo=UTC)
    cell.value = timestamp.astimezone(CHINA_TIME).replace(tzinfo=None, microsecond=0)
    cell.number_format = "yyyy-mm-dd hh:mm:ss"


def _header(sheet, row, labels):
    for column, label in enumerate(labels, 1):
        cell = sheet.cell(row, column)
        _text(cell, label)
        cell.fill = FILL_DARK
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 34


def _sheet_name(owner, used):
    base = re.sub(r"[\[\]:*?/\\\x00-\x1f]", "_", owner or "未分配负责人").strip(" '") or "未分配负责人"
    name = base[:31]
    serial = 2
    while name.casefold() in used:
        suffix = f" ({serial})"
        name = base[:31 - len(suffix)] + suffix
        serial += 1
    used.add(name.casefold())
    return name


def _print_layout(sheet, last_column, last_row):
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_margins = PageMargins(left=.25, right=.25, top=.4, bottom=.4, header=.15, footer=.15)
    sheet.print_options.horizontalCentered = False
    sheet.print_area = f"A1:{last_column}{last_row}"


def _status(product, has_record):
    if not has_record and product.get("status") in {"updating", "failed"}:
        return "更新中（尚无成本记录）" if product["status"] == "updating" else "更新失败（尚无成本记录）"
    return STATUS.get(product.get("status"), "待核算" if not has_record else str(product.get("status") or "待核算"))


def _source(result):
    return "手动" if result and result.get("cost_source") == "manual" else "自动" if result else "待核算"


def build_cost_workbook(rows: list[dict], exported_at: datetime | None = None) -> bytes:
    """Build formula-backed arithmetic from frozen recipe and price inputs."""
    exported_at = exported_at or datetime.now(UTC)
    stamp = exported_at.astimezone(CHINA_TIME).strftime("%Y-%m-%d %H:%M:%S")
    book = Workbook()
    book.calculation.calcMode = "auto"
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    overview = book.active
    overview.title = "总览"
    overview.sheet_view.showGridLines = False
    overview.sheet_view.zoomScale = 85
    overview.merge_cells("A1:K1")
    _text(overview["A1"], "产品成本总览")
    overview["A1"].font = Font(name="Arial", size=16, bold=True, color="233842")
    overview.merge_cells("A2:K2")
    _text(overview["A2"], f"导出时间：{stamp}  ·  在用产品：{len(rows)}  ·  成本单位：元/kg；缺价留空，0 为有效零值")
    overview["A2"].font = Font(name="Arial", size=10, color="52636C")
    _header(overview, 3, ["产品内编", "负责人", "最新优先成本（元/kg）", "库存优先成本（元/kg）", "最新成本来源", "库存成本来源", "收率", "投料条数", "状态", "成本记录时间", "配方版本"])
    for index, item in enumerate(rows, 4):
        product, formula, latest, inventory = (item[key] for key in ("product", "formula", "latest", "inventory"))
        values = [product.get("name"), product.get("owner"), None, None,
                  _source(latest), _source(inventory),
                  None, len(formula.get("lines", [])), _status(product, bool(latest or inventory)), product.get("recorded_at"),
                  f"v{formula.get('revision', product.get('revision', 0))}"]
        for column, value in enumerate(values, 1):
            cell = overview.cell(index, column)
            if column in (3, 4):
                _cost(cell, latest.get("cost") if column == 3 and latest else inventory.get("cost") if inventory else None)
            elif column == 7:
                _number(cell, formula.get("yield"), "0.00%")
            elif column == 8:
                cell.value = value
            elif column == 10:
                _recorded_at(cell, value)
            else:
                _text(cell, value)
            cell.font = Font(name="Arial", size=10, color="233842", bold=column == 1)
            cell.alignment = Alignment(vertical="center", horizontal="right" if column in (3, 4, 7, 8) else "left", indent=1 if column in (2, 5, 6, 9, 10, 11) else 0)
            cell.border = ROW_BORDER
            if index % 2:
                cell.fill = FILL_ALT
        overview.row_dimensions[index].height = 27
    for column, width in {"A": 22, "B": 16, "C": 22, "D": 22, "E": 16, "F": 16, "G": 11, "H": 12, "I": 32, "J": 24, "K": 12}.items():
        overview.column_dimensions[column].width = width
    overview.freeze_panes = "C4"
    overview.auto_filter.ref = f"A3:K{max(3, len(rows) + 3)}"
    _print_layout(overview, "K", max(3, len(rows) + 3))

    by_owner = defaultdict(list)
    for item in rows:
        by_owner[item["product"].get("owner") or "未分配负责人"].append(item)
    used = {"总览"}
    for owner, products in sorted(by_owner.items()):
        sheet = book.create_sheet(_sheet_name(owner, used))
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 85
        sheet.merge_cells("A1:J1")
        _text(sheet["A1"], f"{owner} · 产品配方与成本")
        sheet["A1"].font = Font(name="Arial", size=16, bold=True, color="233842")
        sheet.merge_cells("A2:J2")
        _text(sheet["A2"], f"导出时间：{stamp}  ·  产品：{len(products)}  ·  单价为两种口径实际采用值，回退价格见取价依据")
        sheet["A2"].font = Font(name="Arial", size=10, color="52636C")
        sheet.freeze_panes = "C3"
        for column, width in {"A": 8, "B": 23, "C": 16, "D": 19, "E": 16, "F": 22, "G": 18, "H": 16, "I": 22, "J": 18}.items():
            sheet.column_dimensions[column].width = width
        row = 4
        for item in products:
            product, formula, latest, inventory = (item[key] for key in ("product", "formula", "latest", "inventory"))
            lines = formula.get("lines", [])
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
            title = sheet.cell(row, 1)
            _text(title, f"{product.get('name', '')}  ·  配方 v{formula.get('revision', product.get('revision', 0))}")
            title.fill = FILL_LIGHT
            title.font = Font(name="Arial", size=12, bold=True, color="233842")
            title.alignment = Alignment(vertical="center", indent=1)
            sheet.row_dimensions[row].height = 32
            row += 1
            total = latest.get("total_input") if latest else inventory.get("total_input") if inventory else sum(Decimal(str(line["quantity"])) for line in lines)
            sheet.merge_cells(start_row=row, start_column=6, end_row=row, end_column=7)
            sheet.merge_cells(start_row=row, start_column=9, end_row=row, end_column=10)
            for column in range(1, 11):
                cell = sheet.cell(row, column)
                cell.fill = FILL_ALT
                cell.border = ROW_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            for column, label in ((1, "收率"), (3, "总投料（kg）"), (5, "状态"), (8, "成本记录时间")):
                cell = sheet.cell(row, column)
                _text(cell, label)
                cell.font = Font(name="Arial", size=9, color="52636C")
                cell.alignment = Alignment(vertical="center", indent=1 if column > 1 else 0)
            _number(sheet.cell(row, 2), formula.get("yield"), "0.00%")
            first_line = row + 3
            last_line = first_line + len(lines) - 1
            if lines:
                sheet.cell(row, 4).value = f"=SUM(D{first_line}:D{last_line})"
                sheet.cell(row, 4).number_format = NUMBER_FORMAT
            else:
                _number(sheet.cell(row, 4), total)
            _text(sheet.cell(row, 6), _status(product, bool(latest or inventory)))
            if product.get("recorded_at"):
                _recorded_at(sheet.cell(row, 9), product["recorded_at"])
            else:
                _text(sheet.cell(row, 9), "尚无")
            for column in (2, 4, 6, 9):
                sheet.cell(row, column).font = Font(name="Arial", size=10, color="233842", bold=True)
                sheet.cell(row, column).alignment = Alignment(vertical="center", horizontal="right" if column in (2, 4) else "left", indent=1)
            sheet.row_dimensions[row].height = 28
            row += 1
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            sheet.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
            sheet.merge_cells(start_row=row, start_column=8, end_row=row, end_column=9)
            for start, stop, fill in ((1, 4, FILL_GROUP), (5, 7, FILL_LATEST_GROUP), (8, 10, FILL_INVENTORY_GROUP)):
                for column in range(start, stop + 1):
                    sheet.cell(row, column).fill = fill
            for column, label in ((1, "配方与投料"), (5, f"最新优先成本（{_source(latest)}）"), (8, f"库存优先成本（{_source(inventory)}）")):
                cell = sheet.cell(row, column)
                _text(cell, label)
                cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                cell.alignment = Alignment(vertical="center", indent=1)
            for column, result in ((7, latest), (10, inventory)):
                cell = sheet.cell(row, column)
                if result and result.get("cost") is not None:
                    if result.get("cost_source") == "manual":
                        _cost(cell, result["cost"])
                    elif lines:
                        amount_column = "G" if column == 7 else "J"
                        # ponytail: formulas explain saved costs; add price validation if exports become editable calculators.
                        cell.value = (f'=ROUND(SUM({amount_column}{first_line}:{amount_column}{last_line})/'
                                      f'(D{row - 1}*B{row - 1}),2)')
                        cell.number_format = NUMBER_FORMAT
                cell.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
                cell.alignment = Alignment(vertical="center", horizontal="right", indent=1)
            sheet.row_dimensions[row].height = 32
            row += 1
            for column, label in enumerate(("顺序", "投料内编", "配方比例", "实际投料（kg）", "单价（元/kg）", "取价依据", "投料金额（元）",
                                            "单价（元/kg）", "取价依据", "投料金额（元）"), 1):
                cell = sheet.cell(row, column)
                _text(cell, label)
                cell.fill = FILL_LATEST if 5 <= column <= 7 else FILL_INVENTORY if column >= 8 else FILL_LIGHT
                cell.font = Font(name="Arial", size=9, bold=True, color="3D515A")
                cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
                cell.border = ROW_BORDER
            sheet.row_dimensions[row].height = 30
            row += 1
            total_decimal = Decimal(str(total))
            latest_lines = latest.get("lines", []) if latest else []
            inventory_lines = inventory.get("lines", []) if inventory else []
            for index, line in enumerate(lines):
                left = latest_lines[index] if index < len(latest_lines) else {}
                right = inventory_lines[index] if index < len(inventory_lines) else {}
                sheet.cell(row, 1).value = index + 1
                _text(sheet.cell(row, 2), line.get("code"))
                ratio = Decimal(str(line["ratio"])) / 100 if line.get("ratio") is not None else Decimal(str(line["quantity"])) / total_decimal if total_decimal else None
                _number(sheet.cell(row, 3), ratio, "0.00%")
                _number(sheet.cell(row, 4), line.get("quantity"))
                _number(sheet.cell(row, 5), left.get("unit_cost"))
                _text(sheet.cell(row, 6), PRICE_BASIS.get(left.get("basis"), left.get("basis")) if left else "")
                if left.get("amount") is not None:
                    sheet.cell(row, 7).value = f"=D{row}*E{row}"
                    sheet.cell(row, 7).number_format = NUMBER_FORMAT
                _number(sheet.cell(row, 8), right.get("unit_cost"))
                _text(sheet.cell(row, 9), PRICE_BASIS.get(right.get("basis"), right.get("basis")) if right else "")
                if right.get("amount") is not None:
                    sheet.cell(row, 10).value = f"=D{row}*H{row}"
                    sheet.cell(row, 10).number_format = NUMBER_FORMAT
                for column in (5, 6, 7):
                    sheet.cell(row, column).fill = FILL_LATEST
                for column in (8, 9, 10):
                    sheet.cell(row, column).fill = FILL_INVENTORY
                for column in range(1, 11):
                    cell = sheet.cell(row, column)
                    if index % 2 and column <= 4:
                        cell.fill = FILL_ALT
                    cell.font = Font(name="Arial", size=10, color="233842")
                    cell.alignment = Alignment(vertical="center", horizontal="right" if column in (3, 4, 5, 7, 8, 10) else "center" if column == 1 else "left", indent=1 if column in (2, 6, 9) else 0, wrap_text=True)
                    cell.border = ROW_BORDER
                sheet.row_dimensions[row].height = 29
                row += 1
            sheet.row_dimensions[row].height = 14
            row += 1
        _print_layout(sheet, "J", row - 2)
    output = BytesIO()
    book.save(output)
    return output.getvalue()
