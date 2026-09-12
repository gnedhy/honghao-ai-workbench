"""Normalize the RD5 workbook and calculate its explicit recipe graph, without Excel execution."""
from __future__ import annotations

import hashlib
import io
import re
import zipfile
from collections import defaultdict
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from api.procurement_excel import MAX_FILE, text

PRICE_SHEET = "原料单价"
OMITTED_CODES = {"CF011", "CF1065A", "CF130", "CF745", "CF747", "CF854C", "CF235C"}
ZERO_COST_CODES = {"纯水", "加工S-016（偶联剂C）", "加工S-014 （偶联剂A)"}
DEFAULT_COMPOSITE = "composite:CF401B:CF026K+CF020D"


def number(value, label: str, *, positive: bool = False) -> Decimal:
    try:
        if isinstance(value, bool) or len(str(value)) > 64:
            raise ValueError
        result = Decimal(str(value))
        if not result.is_finite() or result < 0 or result > Decimal("1e18") or abs(result.as_tuple().exponent) > 100 or (positive and result == 0):
            raise ValueError
        return result
    except (InvalidOperation, ValueError):
        raise ValueError(f"{label} 不是有效的{'正' if positive else '非负'}数值") from None


def decimal_text(value: Decimal | None) -> str | None:
    if value is None:
        return None
    value = format(value, "f")
    return value.rstrip("0").rstrip(".") if "." in value else value


def formula(cell):
    return getattr(cell.value, "text", cell.value)


def _check_cached(cell, expected: Decimal | None, label: str):
    if cell.data_type == "e":
        raise ValueError(f"{label} 包含 Excel 错误")
    if cell.value is not None and (expected is None or abs(number(cell.value, label) - expected) > Decimal("0.00000001")):
        raise ValueError(f"{label} 缓存值与独立复算不一致")


def parse_workbook(content: bytes) -> dict:
    if len(content) > MAX_FILE:
        raise ValueError("Excel 文件不能超过8MB")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            files = archive.infolist()
            if len(files) > 2000 or sum(f.file_size for f in files) > 64 * 1024 * 1024:
                raise ValueError("Excel 解压后内容过大")
            if any("vbaproject" in f.filename.lower() or "externallinks/" in f.filename.lower() for f in files):
                raise ValueError("不接受宏或外部链接")
        source = load_workbook(io.BytesIO(content), data_only=False, read_only=True, keep_links=False)
        cached = load_workbook(io.BytesIO(content), data_only=True, read_only=True, keep_links=False)
    except (zipfile.BadZipFile, KeyError, OSError) as error:
        raise ValueError("无法读取有效的xlsx文件") from error
    try:
        if not 2 <= len(source.worksheets) <= 10 or source.sheetnames[0] != PRICE_SHEET:
            raise ValueError("需要原料单价和配方及成本工作表")
        # Materialize bounded cells once; repeated access to a read-only sheet reparses its XML.
        cells, values = {}, {}
        for sheet in source:
            if sheet.max_row > 10000 or sheet.max_column > 30:
                raise ValueError("工作表范围过大")
            cells[sheet.title] = {c.coordinate: c for row in sheet for c in row if getattr(c, "coordinate", None)}
        for sheet in cached:
            values[sheet.title] = {c.coordinate: c for row in sheet for c in row if getattr(c, "coordinate", None)}
        return _parse(cells, values, hashlib.sha256(content).hexdigest())
    finally:
        source.close()
        cached.close()


def _parse(sheets: dict, caches: dict, digest: str) -> dict:
    def get(sheet, address):
        cell = sheets[sheet].get(address)
        return formula(cell) if cell is not None else None

    def cached_check(sheet, address, expected):
        cell = caches[sheet].get(address)
        if cell is not None:
            _check_cached(cell, expected, f"{sheet}!{address}")

    headers = [get(PRICE_SHEET, f"{c}1") for c in "ABCDE"]
    if headers != ["编号", "核算价格", "最新价格", "库存价格", "在途价格"]:
        raise ValueError("原料单价列标题不符合约定")

    def price_cell(address, stack=()):
        if address in stack or len(stack) > 30:
            raise ValueError("原料价格存在循环引用")
        val = get(PRICE_SHEET, address)
        if val is None:
            return None
        if isinstance(val, str) and val.startswith("="):
            match = re.fullmatch(r"=([CD]\d+)(?:/(\d+(?:\.\d+)?))?", val)
            if not match:
                raise ValueError(f"{PRICE_SHEET}!{address} 包含不支持的计价公式")
            base = price_cell(match[1], (*stack, address))
            if base is None:
                raise ValueError(f"{PRICE_SHEET}!{address} 引用缺价单元格")
            result = base / number(match[2] or 1, address, positive=True)
            cached_check(PRICE_SHEET, address, result)
            return result
        return number(val, f"{PRICE_SHEET}!{address}")

    materials = {}
    price_rows = sorted(int(a[1:]) for a in sheets[PRICE_SHEET] if re.fullmatch(r"A\d+", a) and int(a[1:]) > 1)
    for row in price_rows:
        code = text(get(PRICE_SHEET, f"A{row}"))
        if not code:
            continue
        if len(code) > 100 or code in materials:
            raise ValueError(f"原料编号无效或重复：{code}")
        if get(PRICE_SHEET, f"E{row}") is not None:
            raise ValueError(f"{code} 含在途价格，请按已确认的两种价格口径核对")
        latest, stock = price_cell(f"C{row}"), price_cell(f"D{row}")
        expected_formula = f'=IF(ISBLANK(C{row}),IF(ISBLANK(E{row}),IF(ISBLANK(D{row}),"check",D{row}),E{row}),IF(ISNUMBER(C{row}),C{row},"check"))'
        if get(PRICE_SHEET, f"B{row}") not in (None, expected_formula):
            raise ValueError(f"{code} 核算价格公式与约定不一致")
        chosen = latest if latest is not None else stock
        if get(PRICE_SHEET, f"B{row}") is not None and chosen is not None:
            cached_check(PRICE_SHEET, f"B{row}", chosen)
        if (latest == 0 or stock == 0) and code not in ZERO_COST_CODES:
            raise ValueError(f"{code} 的零价格尚未确认")
        materials[code] = {"code": code, "unit": "kg", "source_row": row,
            "latest_price": decimal_text(latest), "inventory_price": decimal_text(stock),
            "source_calculation_price": decimal_text(chosen), "zero_cost": code in ZERO_COST_CODES,
            "source_cells": {f"{c}{row}": get(PRICE_SHEET, f"{c}{row}") for c in "ABCDEF"}}

    recipes = []
    names = {}
    for sheet in list(sheets)[1:]:
        if not sheet.endswith("（配方及成本）"):
            raise ValueError(f"未识别的配方工作表：{sheet}")
        if [get(sheet, f"{c}1") for c in "ABCDEFG"] != ["配方名称", "原料名称", "配方比例%", "原料单价（元/kg）", "实际投料（kg）", "收率%", "成本（元/kg）"]:
            raise ValueError(f"{sheet} 列标题不符合约定")
        groups = defaultdict(list)
        for address in sheets[sheet]:
            if re.fullmatch(r"A\d+", address) and int(address[1:]) > 1 and get(sheet, address) is not None:
                row = int(address[1:])
                name = text(get(sheet, address))
                if not name or not get(sheet, f"B{row}"):
                    raise ValueError(f"{sheet}!{address} 缺少配方名称或投料")
                groups[name].append(row)
        for name, rows in groups.items():
            rows.sort()
            start, end = rows[0], rows[-1]
            if name in names or rows != list(range(start, end + 1)):
                raise ValueError(f"配方名称重复或明细不连续：{name}")
            recipe_id = f"recipe:{name}"
            names[name] = (recipe_id, sheet)
            yield_ = number(get(sheet, f"F{start}"), f"{name}收率", positive=True)
            cost_match = re.fullmatch(rf"=SUMPRODUCT\(C{start}:C(\d+),D{start}:D\1\)/F{start}", str(get(sheet, f"G{start}")))
            if not cost_match or not start <= int(cost_match[1]) <= 10000:
                raise ValueError(f"{name} 成本公式与约定不一致")
            cost_end = int(cost_match[1])
            if any(get(sheet, f"{c}{r}") is not None for r in range(end + 1, cost_end + 1) for c in "ABCD"):
                raise ValueError(f"{name} 成本公式越界引用其他数据")
            lines = []
            total = sum(number(get(sheet, f"E{r}"), f"{name}投料量") for r in rows)
            number(total, f"{name}总投料量", positive=True)
            for r in rows:
                amount = number(get(sheet, f"E{r}"), f"{name}投料量")
                if r > cost_end and text(get(sheet, f"B{r}")) not in ZERO_COST_CODES:
                    raise ValueError(f"{name} 成本公式漏计非零成本原料")
                ratio_formula = get(sheet, f"C{r}")
                match = re.fullmatch(rf"=E{r}/SUM\(E\$?(\d+):E\$?(\d+)\)", str(ratio_formula))
                if isinstance(ratio_formula, (int, float)):
                    if abs(number(ratio_formula, name) - amount / total) > Decimal("0.00000001"):
                        raise ValueError(f"{name} 固定配比与实际投料不一致")
                    denominator = total
                elif not match or not 1 <= int(match[1]) <= int(match[2]) <= 10000:
                    raise ValueError(f"{name} 配比公式无法识别")
                else:
                    denominator = sum(number(get(sheet, f"E{x}"), f"{name}配比分母") for x in range(int(match[1]), int(match[2]) + 1))
                if denominator != total:
                    raise ValueError(f"{name} 配比公式引用了不同总投料量")
                cached_check(sheet, f"C{r}", amount / total)
                if r != start and (get(sheet, f"F{r}") is not None or get(sheet, f"G{r}") is not None):
                    raise ValueError(f"{name} 存在额外收率或成本结果")
                lines.append({"code": text(get(sheet, f"B{r}")), "quantity": decimal_text(amount),
                    "source_row": r, "source_cost_included": r <= cost_end,
                    "source_cells": {f"{c}{r}": get(sheet, f"{c}{r}") for c in "ABCDEFGH"}})
            recipes.append({"id": recipe_id, "name": name, "kind": "recipe", "sheet": sheet,
                "owner": sheet.removesuffix("（配方及成本）").removesuffix("负责"), "source_row": start,
                "yield": decimal_text(yield_), "lines": lines, "cached_cost": get_cached(caches, sheet, f"G{start}")})

    composites_by_cell = {}
    for address in sorted(sheets[PRICE_SHEET], key=lambda a: (int(re.search(r"\d+", a)[0]), a)):
        if not re.fullmatch(r"H\d+", address) or get(PRICE_SHEET, address) is None:
            continue
        r = int(address[1:])
        if get(PRICE_SHEET, address) != "CF401B":
            raise ValueError("仅识别已确认的 CF401B 复配区")
        codes = [text(get(PRICE_SHEET, f"I{x}")) for x in [r, r + 1]]
        if codes[0] not in {"CF026", "CF026K"} or codes[1] not in {"CF020C", "CF020D"}:
            raise ValueError("CF401B 复配原料与约定不一致")
        recipe_id = "composite:CF401B:" + "+".join(codes)
        if recipe_id in composites_by_cell.values():
            raise ValueError("复配方案重复")
        amounts = [number(get(PRICE_SHEET, f"L{x}"), "复配投料量", positive=True) for x in [r, r + 1]]
        yield_ = number(get(PRICE_SHEET, f"M{r}"), "复配收率", positive=True)
        if amounts != [Decimal(3000), Decimal(300)] or yield_ != Decimal("0.995"):
            raise ValueError("CF401B 投料或收率与已确认方案不同")
        if get(PRICE_SHEET, f"N{r}") != f"=SUMPRODUCT(J{r}:J{r+1},K{r}:K{r+1})/M{r}":
            raise ValueError("CF401B 成本公式与约定不一致")
        lines = []
        for x, code, amount in zip([r, r + 1], codes, amounts):
            if get(PRICE_SHEET, f"J{x}") != f"=L{x}/SUM(L{r}:L{r+1})" or get(PRICE_SHEET, f"K{x}") != f"=VLOOKUP(I{x},原料单价!A:B,2,0)":
                raise ValueError("CF401B 配比或单价引用无法识别")
            lines.append({"code": code, "kind": "material", "ref": code, "quantity": decimal_text(amount),
                "source_row": x, "source_cells": {f"{c}{x}": get(PRICE_SHEET, f"{c}{x}") for c in "HIJKLMN"}})
        composites_by_cell[f"N{r}"] = recipe_id
        recipes.append({"id": recipe_id, "name": "CF401B（" + "+".join(codes) + "）", "kind": "composite",
            "sheet": PRICE_SHEET, "owner": "研发五部", "source_row": r, "yield": decimal_text(yield_),
            "lines": lines, "cached_cost": get_cached(caches, PRICE_SHEET, f"N{r}")})
    if len(composites_by_cell) != 4 or DEFAULT_COMPOSITE not in composites_by_cell.values():
        raise ValueError("需要四种完整的 CF401B 复配方案")

    for recipe in recipes:
        if recipe["kind"] == "composite":
            continue
        for line in recipe["lines"]:
            code, r = line["code"], line["source_row"]
            raw = line["source_cells"][f"D{r}"]
            if raw == f"=VLOOKUP(B{r},原料单价!A:B,2,0)":
                line.update(kind="material", ref=code)
            elif isinstance(raw, (int, float)):
                line.update(kind="material", ref=code, source_price_override=decimal_text(number(raw, code)))
            elif isinstance(raw, str) and raw.startswith("=原料单价!") and raw[6:] in composites_by_cell:
                line.update(kind="composite", ref=composites_by_cell[raw[6:]])
            elif code in names and raw in (f"=VLOOKUP(B{r},A:G,7,0)", f"=VLOOKUP(B{r},'{names[code][1]}'!A:G,7,0)"):
                if raw == f"=VLOOKUP(B{r},A:G,7,0)" and names[code][1] != recipe["sheet"]:
                    raise ValueError(f"{code} 配方引用工作表不匹配")
                line.update(kind="recipe", ref=names[code][0])
            else:
                raise ValueError(f"{recipe['sheet']}!D{r} 单价引用无法识别：{raw}")
    for recipe in recipes:
        for line in recipe["lines"]:
            if line["kind"] == "material" and line["code"] not in materials:
                raise ValueError(f"{line['code']} 缺少原料单价记录")
    package = {"format_version": 1, "source_sha256": digest, "department": "研发五部",
        "price_policy": "latest_then_inventory", "default_composite": DEFAULT_COMPOSITE,
        "materials": list(materials.values()), "recipes": recipes,
        "notes": ["CF007B、CF063A保留原表计价关联，当前核算取台账价格。", "RH-9收率102%沿用原表，待业务核对。"]}
    audit = calculate(package, {}, source_mode=True)
    for recipe in recipes:
        cost = audit[recipe["id"]]["cost"]
        if recipe["cached_cost"] is not None:
            if cost is None or abs(number(recipe["cached_cost"], recipe["name"]) - Decimal(cost)) > Decimal("0.00000001"):
                raise ValueError(f"{recipe['name']} 原表成本缓存与复算不一致")
    package["source_replay"] = audit
    return package


def get_cached(caches, sheet, address):
    cell = caches[sheet].get(address)
    if cell is None or cell.value is None:
        return None
    return decimal_text(number(cell.value, f"{sheet}!{address}"))


def calculate(package: dict, prices: dict, *, source_mode: bool = False, price_policy: str = "latest") -> dict:
    if price_policy not in {"latest", "inventory"}:
        raise ValueError("无效取价口径")
    recipes = {r["id"]: r for r in package["recipes"]}
    if len(recipes) != len(package["recipes"]):
        raise ValueError("配方标识重复")
    source_materials = {m["code"]: m for m in package["materials"]}
    policy = {} if source_mode else package.get("current_policy", {})
    historical = set(policy.get("historical_recipe_ids", []))
    if historical - recipes.keys():
        raise ValueError("历史试算标识不存在")
    results = {}

    def visit(key, stack=()):
        if key in stack or len(stack) > 100:
            raise ValueError("配方循环引用或层级过深：" + " → ".join((*stack, key)))
        if key in results:
            return results[key]
        if key not in recipes:
            raise ValueError(f"未找到引用的配方：{key}")
        if key in historical:
            raise ValueError(f"当前配方仍引用历史试算：{key}")
        recipe = recipes[key]
        yield_ = number(recipe["yield"], recipe["name"] + "收率", positive=True)
        amount = Decimal(0)
        weighted = Decimal(0)
        missing, lines = set(), []
        for line in recipe["lines"]:
            quantity = number(line["quantity"], recipe["name"] + "投料量")
            amount += quantity
            kind, ref = line["kind"], line["ref"]
            if kind == "material":
                ref = policy.get("material_replacements", {}).get(ref, ref)
            if not source_mode and kind == "material" and ref == "CF401B":
                kind, ref = "composite", package["default_composite"]
            absent = []
            if kind != "material":
                nested = visit(ref, (*stack, key))
                price = Decimal(nested["cost"]) if nested["cost"] is not None else None
                absent = nested["missing_materials"]
                basis = kind
            elif source_mode:
                price_text = line.get("source_price_override", source_materials[ref]["source_calculation_price"])
                # Excel VLOOKUP returns 0 for an empty result cell. Reproduce it only in the labeled source audit.
                price = Decimal(price_text) if price_text is not None else Decimal(0)
                basis = "source_blank_as_zero" if price_text is None else "source_workbook"
            else:
                entry = prices.get(ref, {})
                latest, stock = entry.get("latest_price"), entry.get("inventory_price")
                stock = stock if stock is not None and number(stock, ref) > 0 else None
                basis = "inventory" if price_policy == "inventory" and stock is not None else "latest" if latest is not None else "inventory" if stock is not None else "missing"
                selected = {"inventory": stock, "latest": latest, "missing": None}[basis]
                price = number(selected, ref) if selected is not None else None
                absent = [ref] if price is None else []
                basis = entry.get(basis + "_basis", basis)
            if quantity > 0:
                missing.update(absent)
                if price is not None and (not source_mode or line.get("source_cost_included", True)):
                    weighted += quantity * price
            lines.append({"source_row": line["source_row"], "code": line["code"], "kind": kind, "ref": ref,
                "quantity": line["quantity"], "unit_cost": decimal_text(price), "basis": basis,
                "amount": decimal_text(quantity * price) if price is not None else None, "missing_materials": absent})
        number(amount, recipe["name"] + "总投料量", positive=True)
        cost = None if missing else weighted / amount / yield_
        result = {"cost": decimal_text(cost), "missing_materials": sorted(missing), "total_input": decimal_text(amount),
            "output_quantity": decimal_text(amount * yield_), "yield": recipe["yield"], "lines": lines}
        results[key] = result
        return result

    for key in recipes:
        if key not in historical:
            visit(key)
    return results
