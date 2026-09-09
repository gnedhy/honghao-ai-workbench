"""Bounded, local XLSX extraction and atomic initial historical migration."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sqlite3
import zipfile
from collections import Counter
from datetime import UTC, date, datetime
from decimal import Decimal, DecimalException
from xml.etree.ElementTree import ParseError
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

TOTAL = "原料行情总表"
DEPARTMENTS = ["研发一部", "研发二部", "研发三部", "研发四部", "研发五部", "宏昊生物"]
MAX_FILE = 8 * 1024 * 1024


def read_workbook(content: bytes) -> dict:
    if len(content) > MAX_FILE:
        raise ValueError("Excel 文件不能超过8MB")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 2000 or sum(i.file_size for i in infos) > 64 * 1024 * 1024:
                raise ValueError("Excel 解压后内容过大")
            if any("vbaproject" in i.filename.lower() or "externallinks/" in i.filename.lower() for i in infos):
                raise ValueError("不接受宏或外部链接")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        result = {}
        try:
            if len(workbook.worksheets) > 30:
                raise ValueError("工作表数量过多")
            for sheet in workbook:
                if sheet.max_row > 10000 or sheet.max_column > 100:
                    raise ValueError("工作表范围过大")
                rows = []
                for row in sheet.iter_rows():
                    if any(cell.data_type == "f" for cell in row):
                        raise ValueError("请将公式转换为值后导入")
                    rows.append([cell.value for cell in row])
                if rows:
                    result[sheet.title] = rows
        finally:
            workbook.close()
        return result
    except (zipfile.BadZipFile, KeyError, OSError, ParseError, TypeError, IndexError) as error:
        raise ValueError("无法读取有效的xlsx文件") from error


def text(value) -> str:
    return "" if value is None else str(value).strip()


def day(value) -> str | None:
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    match = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", text(value))
    if match:
        try:
            return date(*map(int, match.groups())).isoformat()
        except ValueError:
            return None
    return None


def price_value(value) -> tuple[str | None, str]:
    raw = text(value)
    if not raw:
        return None, "missing"
    try:
        number = Decimal(raw)
        if number.is_finite() and number >= 0 and len(raw) <= 64 and -8 <= number.as_tuple().exponent <= 18 and number.adjusted() < 18:
            rendered = format(number, "f")
            return (rendered.rstrip("0").rstrip(".") if "." in rendered else rendered), "number"
    except DecimalException:
        pass
    return None, "range" if re.fullmatch(r"\d+(?:\.\d+)?\s*[-~～]\s*\d+(?:\.\d+)?", raw) else "invalid"


def inspect_workbook(content: bytes) -> dict:
    sheets = read_workbook(content)
    return {"sheets": [{"name": name, "dates": [d for v in rows[0] if (d := day(v))], "rows": len(rows)-1} for name, rows in sheets.items()], "default_sheet": TOTAL if TOTAL in sheets else next(iter(sheets), "")}


def selected_prices(content: bytes, sheet: str, price_date: str, known: dict) -> dict:
    sheets = read_workbook(content)
    if sheet not in sheets:
        raise ValueError("工作表不存在")
    rows = sheets[sheet]
    dates = [day(v) for v in rows[0]]
    if price_date not in dates:
        raise ValueError("所选日期不存在，请核对表头")
    column = dates.index(price_date)
    codes = Counter(text(r[1]) for r in rows[1:] if len(r)>1 and text(r[1]))
    result = []
    csv_rows = []
    for n, row in enumerate(rows[1:], 2):
        code = text(row[1]) if len(row)>1 else ""
        if not code:
            continue
        raw = text(row[column])
        value, kind = price_value(raw)
        state = "duplicate" if codes[code]>1 else "unknown" if code not in known else kind
        result.append({"row": n, "code": code, "raw_price": raw, "price": value, "state": state})
        if state == "number":
            item = known[code]
            csv_rows.append([code, item["name"], item["unit"], value])
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(["编号", "名称", "单位", "最新价"])
    writer.writerows(csv_rows)
    return {"rows": result, "content": stream.getvalue(), "valid": len(csv_rows),
            "blocked": sum(r["state"] in {"duplicate", "unknown", "range", "invalid"} for r in result),
            "unreported": sum(r["state"] == "missing" for r in result)}


def history_preview(content: bytes) -> dict:
    sheets = read_workbook(content)
    if TOTAL not in sheets or not all(name in sheets for name in DEPARTMENTS):
        raise ValueError("首次迁入需要原料行情总表及六个部门子表")
    rows = sheets[TOTAL]
    dates = [(i, d) for i,v in enumerate(rows[0]) if (d := day(v))]
    total_codes = [text(r[1]) for r in rows[1:] if len(r)>1 and text(r[1])]
    if len(set(total_codes)) != len(total_codes) or len(set(d for _,d in dates)) != len(dates):
        raise ValueError("总表原料编号或日期重复")
    codes = {text(r[1]) for name in [TOTAL, *DEPARTMENTS] for r in sheets[name][1:] if len(r)>1 and text(r[1])}
    anomalies = [{"sheet": TOTAL, "cell": f"{get_column_letter(i+1)}{n}", "code": text(row[1]), "date": d, "value": text(row[i]), "kind": price_value(row[i])[1]}
                 for n,row in enumerate(rows[1:],2) if text(row[1]) for i,d in dates if price_value(row[i])[1] in {"range","invalid"}]
    return {"material_count": len(codes), "history_material_count": len(total_codes), "dates": sorted(d for _,d in dates),
            "departments": [{"name": name, "count": len({text(r[1]) for r in sheets[name][1:] if text(r[1])})} for name in DEPARTMENTS],
            "anomalies": anomalies, "sha256": hashlib.sha256(content).hexdigest()}


def import_history(path: Path, source: Path, actor_id: str, *, expected_sha256: str) -> dict:
    content = source.read_bytes()
    preview = history_preview(content)
    if preview["sha256"] != expected_sha256:
        raise ValueError("文件已变化，请重新预检")
    sheets = read_workbook(content)
    digest = preview["sha256"]
    now = datetime.now(UTC).isoformat()
    stored = path.parent / "controlled-work" / "procurement-sources" / f"{digest}.xlsx"
    created_archive = False
    try:
        with sqlite3.connect(path) as db:
            db.execute("BEGIN IMMEDIATE")
            admin = db.execute("SELECT access_level,is_active FROM identity_users WHERE id=?", (actor_id,)).fetchone()
            if not admin or admin != (5,1):
                raise PermissionError("首次迁入仅限有效管理员")
            if db.execute("SELECT 1 FROM procurement_source_imports WHERE sha256=?", (digest,)).fetchone():
                raise ValueError("这份文件已迁入，不可重复执行")
            if db.execute("SELECT 1 FROM procurement_materials LIMIT 1").fetchone():
                raise ValueError("首次迁入必须使用空采购库，不覆盖现有业务数据")
            stored.parent.mkdir(parents=True, exist_ok=True)
            if stored.exists() and stored.read_bytes() != content:
                raise ValueError("来源归档校验失败")
            if not stored.exists():
                with stored.open("xb") as stream:
                    created_archive = True
                    stream.write(content)
            db.execute("INSERT INTO procurement_source_imports VALUES (?,?,?,?,?)", (digest, source.name, str(stored), actor_id, now))
            ids = {}
            for name in [TOTAL, *DEPARTMENTS]:
                for n,row in enumerate(sheets[name][1:],2):
                    code = text(row[1]) if len(row)>1 else ""
                    if not code:
                        continue
                    if code not in ids:
                        ids[code] = str(uuid4())
                        db.execute("INSERT INTO procurement_materials(id,code,name,unit,updated_at) VALUES (?,?,?,'kg',?)", (ids[code], code, code, now))
                    db.execute("INSERT INTO procurement_material_sources VALUES (?,?,?,?,?)", (ids[code],name,n,json.dumps(row,ensure_ascii=False,default=str),digest))
            for position,name in enumerate(DEPARTMENTS):
                group = str(uuid4())
                db.execute("INSERT INTO procurement_departments VALUES (?,?,?)", (group,name,position))
                db.executemany("INSERT INTO procurement_department_materials VALUES (?,?)", [(group,ids[code]) for code in {text(r[1]) for r in sheets[name][1:] if text(r[1])}])
            rows = sheets[TOTAL]
            columns = sorted([(i, day(v)) for i,v in enumerate(rows[0]) if day(v)], key=lambda item:item[1])
            previous = {}
            prior_batch = None
            for version,(column,price_date) in enumerate(columns,1):
                batch = str(uuid4())
                reported = 0
                for n,row in enumerate(rows[1:],2):
                    code = text(row[1])
                    if not code:
                        continue
                    raw = text(row[column])
                    if raw:
                        reported += 1
                        value,kind = price_value(raw)
                        previous[code] = (value,kind,price_date,raw,f"{get_column_letter(column+1)}{n}")
                    value,kind,origin_day,raw,cell = previous.get(code,(None,"missing",None,None,None))
                    db.execute("INSERT INTO procurement_price_batch_items VALUES (?,?,?,?,'kg',?,NULL,NULL,?)", (batch,ids[code],code,code,value,value))
                    db.execute("INSERT INTO procurement_snapshot_sources VALUES (?,?,?,?,?,?,?,?)", (batch,ids[code],origin_day,raw,kind,TOTAL,cell,digest))
                db.execute("INSERT INTO procurement_price_batches(id,version,published_by,published_at,item_count,price_date,source_name,base_batch_id) VALUES (?,?,?, ?,?,?,?,?)",
                           (batch,version,actor_id,now,preview["history_material_count"],price_date,source.name,prior_batch))
                db.execute("INSERT INTO procurement_batch_provenance VALUES (?,'historical_import',?,?,?,?)", (batch,digest,now,actor_id,reported))
                prior_batch = batch
            for code,(value,kind,origin_day,raw,cell) in previous.items():
                db.execute("UPDATE procurement_materials SET latest_price=? WHERE id=?", (value,ids[code]))
            # Only source-attributed values are stored; no invented historical user actions.
            from api.procurement_collaboration import admin_event
            admin_event(db,actor_id,"history.imported",digest,preview)
    except Exception:
        if created_archive and stored.is_file():
            stored.unlink()
        raise
    return preview
