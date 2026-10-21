"""Department projection does not leak recursively frozen formulas or bypass writes."""
from copy import deepcopy
from io import BytesIO
import json
import pytest
from openpyxl import load_workbook

from api.identity import IdentityStore
from api.postgres import transaction
from api.research import capture,enqueue,packed
from api.procurement_collaboration import set_grant
from tests.helpers import TEST_ADMIN_PASSWORD
from tests.test_research_workbench import ready,K,body_for,records
from tests.test_procurement_rd5 import data,trial_data

PREFIX='/api/workbenches/research'
FOREIGN='recipe:department-other'

@pytest.fixture
def departments(ready):
    settings,client,admin,store=ready
    identity=IdentityStore(store.url)
    user=identity.create_user(username='rd1',display_name='本部门负责人',department='研发一部',password=TEST_ADMIN_PASSWORD,scope_levels={'research':3})
    other=identity.create_user(username='rd2',display_name='SECRET-OWNER',department='研发二部',password=TEST_ADMIN_PASSWORD,scope_levels={'research':3})
    root=identity.get_user(admin)
    for person in (user,other):set_grant(store.url,root,person['id'],True,False,'research')
    with transaction(store.url,write=True) as db:
        original=next(f for f in capture(db)['package']['recipes'] if f['id']==K)
        code=next(c for c,v in capture(db)['prices'].items() if v['latest_price'] is not None and v['inventory_price'] is not None)
        foreign=dict(deepcopy(original),id=FOREIGN,name='OTHER-PRODUCT',owner='SECRET-OWNER',source_kind='erp_bom',department='研发二部',department_id=other['primary_department_id'],lines=[dict(kind='material',ref=code,code=code,quantity='1',ratio='777.123',source_row=91234,source_cells={'private':'SECRET-FORMULA'})])
        original.update(source_kind='erp_bom',department='研发一部',department_id=user['primary_department_id'],owner=user['display_name'],lines=[dict(kind='recipe',ref=FOREIGN,code=foreign['name'],quantity='1',source_row=991)],**{'yield':'1'})
        db.execute('INSERT INTO research_formulas(id,position,formula,revision) VALUES(%s,999,%s,1)',(FOREIGN,packed(foreign)))
        db.execute('UPDATE research_formulas SET formula=%s WHERE id=%s',(packed(original),K))
        downstream=dict(deepcopy(foreign),id='recipe:department-downstream',name='SECRET-DOWNSTREAM',lines=[dict(kind='recipe',ref=K,code=original['name'],quantity='1',source_row=993)])
        db.execute('INSERT INTO research_formulas(id,position,formula,revision) VALUES(%s,1000,%s,1)',(downstream['id'],packed(downstream)))
        enqueue(db,'部门隔离合成验收')
    store.process_events()
    client.post('/api/logout')
    assert client.post('/api/login',json={'username':'rd1','password':TEST_ADMIN_PASSWORD}).status_code==200
    return settings,client,admin,store,user,other

def test_department_lists_history_details_export_and_reference_cost(departments):
    _,client,_,store,user,_=departments
    from api.research_access import ResearchAccess
    source={"latest":{"lines":[{"kind":"recipe","ref":FOREIGN,"code":"OTHER-PRODUCT","missing_materials":["SECRET-MATERIAL"]}]},"graph":[{"id":FOREIGN}]}
    original=deepcopy(source)
    with transaction(store.url) as db:
        projected=ResearchAccess(db,user).project(source,{K})
    assert source==original and projected["graph"]==[]
    assert projected["latest"]["lines"][0]["missing_materials"]==["OTHER-PRODUCT"]
    frozen=records(store.url)
    listing=client.get(PREFIX+'/products').json()
    assert [p['id'] for p in listing['products']]==[K]
    assert client.get(PREFIX+'/products/'+FOREIGN).status_code==404
    for route in ('/overview','/formulas','/history','/trials'):
        response=client.get(PREFIX+route)
        assert response.status_code==200,response.text
        assert 'SECRET-' not in response.text
    detail=client.get(PREFIX+'/products/'+K)
    assert detail.status_code==200,detail.text
    assert 'SECRET-' not in detail.text and '777.123' not in detail.text
    data=detail.json()
    assert all(f['id']!=FOREIGN for f in data['recipes'])
    assert data['latest']['lines'][0]['unit_cost'] is not None
    assert data['latest']['lines'][0]['restricted_reference'] is True
    assert data['options']['owners']==[user['display_name']]
    for version in client.get(PREFIX+'/history').json()['versions']:
        result=client.get(PREFIX+'/history/'+version['id'])
        assert result.status_code==200 and 'SECRET-' not in result.text and '777.123' not in result.text
        summary=client.get(PREFIX+'/history/'+version['id'],params={'summary':True})
        assert summary.status_code==200 and 'SECRET-' not in summary.text
        assert all('formula' not in row and 'graph' not in row for row in summary.json()['products'])
        focused=client.get(PREFIX+'/history/'+version['id'],params={'product_id':K})
        assert focused.status_code==200 and len(focused.json()['products'])==1
        assert focused.json()['products'][0]==next(row for row in result.json()['products'] if row['id']==K)
        assert client.get(PREFIX+'/history/'+version['id'],params={'product_id':FOREIGN}).status_code==404
    book=load_workbook(BytesIO(client.get(PREFIX+'/products/export').content),data_only=False)
    assert book['总览'].max_row==4
    assert all('SECRET-' not in str(cell.value) and '777.123' not in str(cell.value) for ws in book for row in ws for cell in row)
    book.close()
    assert records(store.url)==frozen

def test_foreign_writes_and_membership_revocation_are_denied(departments):
    _,client,_,store,user,_=departments
    frozen=records(store.url)
    body=body_for(store,FOREIGN)
    for method,path,payload in [('POST','/simulate',body),('PUT','/draft',dict(body,simulation_token='invalid')),('DELETE','/draft',{'revision':1}),('POST','/activate',{'revision':1,'simulation_token':'invalid'}),('DELETE','',{'revision':0})]:
        response=client.request(method,PREFIX+'/products/'+FOREIGN+path,json=payload)
        assert response.status_code==404,response.text
    own=body_for(store)
    trial=client.post(PREFIX+'/products/'+K+'/simulate',json=own)
    assert trial.status_code==200 and 'SECRET-' not in trial.text
    saved=client.put(PREFIX+'/products/'+K+'/draft',json=dict(own,simulation_token=trial.json()['simulation_token']))
    assert saved.status_code==200,saved.text
    persisted=client.get(PREFIX+'/products/'+K)
    assert persisted.status_code==200 and 'SECRET-' not in persisted.text
    assert persisted.json()['draft']['simulation']['affected']==[{'id':K,'name':store.detail(K)['product']['name']}]
    with transaction(store.url,write=True) as db:db.execute('DELETE FROM organization_memberships WHERE user_id=%s',(user['id'],))
    assert client.get(PREFIX+'/products').json()['products']==[]
    assert client.get(PREFIX+'/products/'+K).status_code==404
    activated=client.post(PREFIX+'/products/'+K+'/activate',json=saved.json())
    assert activated.status_code==404,activated.text
    assert records(store.url)==frozen

def test_same_department_activation_and_cross_department_grant_boundary(departments):
    _,client,admin,store,user,other=departments
    original=records(store.url)
    body=body_for(store)
    trial=client.post(PREFIX+'/products/'+K+'/simulate',json=body).json()
    saved=client.put(PREFIX+'/products/'+K+'/draft',json=dict(body,simulation_token=trial['simulation_token']))
    result=client.post(PREFIX+'/products/'+K+'/activate',json=saved.json())
    assert result.status_code==200,result.text
    store.process_events()
    assert records(store.url)[:len(original)]==original
    root=IdentityStore(store.url).get_user(admin)
    set_grant(store.url,root,user['id'],True,True,'research')
    grant_list=client.get(PREFIX+'/activation-grants')
    assert grant_list.status_code==200 and 'SECRET-' not in grant_list.text
    denied=client.put(PREFIX+'/activation-grants/'+other['id'],json={'enabled':False})
    assert denied.status_code==403,denied.text

def test_new_formula_uses_membership_and_rejects_forged_department(departments):
    _,client,_,store,user,other=departments
    refused=client.post(PREFIX+'/formulas',json={'name':'FORGED','owner':user['display_name'],'department_id':other['primary_department_id']})
    assert refused.status_code==403,refused.text
    created=client.post(PREFIX+'/formulas',json={'name':'OWN-NEW','owner':user['display_name']})
    assert created.status_code==200,created.text
    assert created.json()['department_id']==user['primary_department_id']
    assert client.get(PREFIX+'/products/'+created.json()['id']).status_code==200

def test_missing_upstream_material_is_not_exposed_in_history_summary(departments):
    _,client,_,store,_,_=departments
    with transaction(store.url,write=True) as db:
        foreign=json.loads(db.execute('SELECT formula FROM research_formulas WHERE id=%s',(FOREIGN,)).fetchone()[0])
        foreign['lines'][0].update(code='SECRET-MATERIAL',ref='SECRET-MATERIAL')
        db.execute('UPDATE research_formulas SET formula=%s WHERE id=%s',(packed(foreign),FOREIGN))
        enqueue(db,'核对上游缺价投影')
    store.process_events()
    for route in ('/products','/history','/overview','/products/'+K):
        response=client.get(PREFIX+route)
        assert response.status_code==200,response.text
        assert 'SECRET-' not in response.text
    row=client.get(PREFIX+'/products').json()['products'][0]
    assert row['status']=='missing' and row['missing_materials']==['OTHER-PRODUCT']

def test_approved_accounts_are_dry_run_until_import_and_keep_existing_passwords(ready):
    from scripts.configure_research_accounts import configure,PEOPLE,OFFICE_PEOPLE,RESEARCH_DEPARTMENTS
    from api.organization import legacy_department
    settings,_,admin,store=ready
    identities=IdentityStore(store.url)
    existing=identities.create_user(username='existing-activator',display_name='李树荣',department='研发一部',password='Keep-Existing-2026',scope_levels={'research':2})
    office=[identities.create_user(username='office-'+str(i),display_name=name,department='总经办',password='Keep-Office-2026',scope_levels={'research':4 if i==1 else 2,'procurement':2}) for i,name in enumerate(OFFICE_PEOPLE)]
    set_grant(store.url,identities.get_user(admin),office[1]['id'],True,True,'research')
    with transaction(store.url) as db:
        grants=db.execute('SELECT * FROM research_activation_grants WHERE user_id=ANY(%s)',([u['id'] for u in office],)).fetchall()
    before=len(identities.list_users())
    plan=configure(store.url,admin)
    assert len(plan)==9 and len(identities.list_users())==before
    with pytest.raises(ValueError,match='统一导入'):configure(store.url,admin,True)
    assert len(identities.list_users())==before
    with transaction(store.url,write=True) as db:
        for name in RESEARCH_DEPARTMENTS:legacy_department(db,name)
        package=json.loads(db.execute('SELECT payload FROM research_meta WHERE id=1').fetchone()[0]);package['erp_import']={'inputs_sha256':'synthetic-verified-import'}
        db.execute('UPDATE research_meta SET payload=%s WHERE id=1',(packed(package),))
    configure(store.url,admin,True)
    assert identities.login('existing-activator','Keep-Existing-2026',3600)
    assert len(identities.list_users())==before+4
    for _,name,dept in PEOPLE:
        user=next(u for u in identities.list_users() if u['display_name']==name)
        assert user['department']==dept and user['scope_levels']['research']==3
        with transaction(store.url) as db:assert db.execute('SELECT manager FROM research_activation_grants WHERE user_id=%s',(user['id'],)).fetchone()==(0,)
    for original in office:
        user=identities.get_user(original['id'])
        assert user['primary_department_id']==original['primary_department_id'] and user['department']=='总经办'
        assert user['scope_levels']==original['scope_levels'] and identities.login(user['username'],'Keep-Office-2026',3600)
        assert RESEARCH_DEPARTMENTS <= {d['name'] for d in user['departments']}
    with transaction(store.url) as db:
        assert db.execute('SELECT * FROM research_activation_grants WHERE user_id=ANY(%s)',([u['id'] for u in office],)).fetchall()==grants


def test_detail_does_not_read_unrelated_product_history(departments,monkeypatch):
    _,_,_,store,user,_=departments
    with transaction(store.url,write=True) as db:
        other=json.loads(db.execute('SELECT formula FROM research_formulas WHERE id=%s',(K,)).fetchone()[0])
        other.update(id='recipe:own-unrelated',name='OWN-UNRELATED')
        db.execute('INSERT INTO research_formulas(id,position,formula,revision) VALUES(%s,9000,%s,1)',(other['id'],packed(other)))
    seen=[]
    original=store._linked_history
    def history(db,key,*args):
        seen.append(key)
        return original(db,key,*args)
    monkeypatch.setattr(store,'_linked_history',history)
    detail=store.detail(K,user)
    assert set(seen)=={K}
    assert [r['id'] for r in detail['recipes']]==[K]
    assert detail['product']['latest_cost']==next(p for p in store.listing(user)['products'] if p['id']==K)['latest_cost']


def test_history_reads_are_batched_and_preserve_frozen_comparisons(ready,monkeypatch):
    _,_,_,store=ready
    from api.research_history import linked_history
    frozen=records(store.url)
    with transaction(store.url) as db:
        keys=[r[0] for r in db.execute('SELECT id FROM research_formulas')]
        expected={key:linked_history(db,key) for key in keys}
        connection_type=type(db)
    queries=[]
    execute=connection_type.execute
    def counted(self,query,*args,**kwargs):
        if isinstance(query,str) and query.startswith('SELECT product_id,') and ('FROM research_cost_records' in query or 'FROM research_backfill_records' in query):
            queries.append(query)
        return execute(self,query,*args,**kwargs)
    monkeypatch.setattr(connection_type,'execute',counted)
    rows=store.listing()['products']
    for row in rows:
        history=expected[row['id']]
        assert row['change']==history[0]['change']
        assert row['comparison_basis']==history[0]['comparison_basis']
    assert len(queries)==2
    for version in store.history()['versions']:
        for row in version['products']:
            original=next(r for r in expected[row['id']] if r['event_id']==row['event_id'])
            assert row['change']==original['change'] and row['comparison_basis']==original['comparison_basis']
    assert len(queries)==4 and records(store.url)==frozen


def test_overview_uses_one_snapshot_across_concurrent_catalog_change(departments,monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    _,_,_,store,user,_=departments
    original=store.listing
    def listing(actor=None,histories=None):
        result=original(actor,histories)
        def change():
            with transaction(store.url,write=True) as db:
                formula=json.loads(db.execute('SELECT formula FROM research_formulas WHERE id=%s',(K,)).fetchone()[0])
                formula['owner']='CHANGED-AFTER-SNAPSHOT'
                db.execute('UPDATE research_formulas SET formula=%s WHERE id=%s',(packed(formula),K))
        with ThreadPoolExecutor(max_workers=1) as pool:pool.submit(change).result(timeout=10)
        return result
    monkeypatch.setattr(store,'listing',listing)
    snapshot=store.overview(user)
    assert snapshot['products'][0]['owner']==user['display_name']
    assert next(f for f in snapshot['formulas'] if f['id']==K)['owner']==user['display_name']
    assert 'CHANGED-AFTER-SNAPSHOT' not in snapshot['owners']
    assert next(f for f in store.formulas(user)['formulas'] if f['id']==K)['owner']=='CHANGED-AFTER-SNAPSHOT'


def test_overview_reuses_history_and_returns_counts_without_losing_costs(ready,monkeypatch):
    _,_,_,store=ready
    frozen=records(store.url)
    expected={**store.listing(),**store.history(),**store.formulas()}
    for row in expected['products']+expected['formulas']:
        if isinstance(row.get('lines'),list):row['lines']=len(row['lines'])
    with transaction(store.url) as db:connection_type=type(db)
    execute=connection_type.execute
    repeated=[]
    def counted(self,query,*args,**kwargs):
        if isinstance(query,str) and query.startswith('SELECT DISTINCT ON (product_id)'):repeated.append(query)
        return execute(self,query,*args,**kwargs)
    monkeypatch.setattr(connection_type,'execute',counted)
    assert store.overview()==expected
    assert repeated==[] and records(store.url)==frozen
    assert len(store.detail(K)['recipes'][0]['lines'])>0


def test_stored_summary_tracks_payload_and_cannot_be_written_independently(ready):
    _,_,_,store=ready
    projection="(payload::jsonb - ARRAY['graph','calculations','formula','source_hashes'])::text"
    with transaction(store.url) as db:
        for table in ('research_cost_records','research_backfill_records'):
            expected=projection if table=='research_cost_records' else "(payload::jsonb - ARRAY['graph','calculations','formula','source_hashes','latest','inventory','lines'])::text"
            assert db.execute(f'SELECT COUNT(*) FROM {table} WHERE summary_payload IS DISTINCT FROM {expected}').fetchone()[0]==0
        sequence,raw=db.execute('SELECT sequence,payload FROM research_cost_records ORDER BY sequence DESC LIMIT 1').fetchone()
    updated=json.loads(raw);updated['latest_cost']='0.123456789012345678';updated['inventory_cost']=None
    with transaction(store.url,write=True) as db:
        db.execute('UPDATE research_cost_records SET payload=%s WHERE sequence=%s',(packed(updated),sequence))
        assert db.execute(f'SELECT summary_payload={projection} FROM research_cost_records WHERE sequence=%s',(sequence,)).fetchone()[0]
    import psycopg
    with pytest.raises(psycopg.errors.GeneratedAlways):
        with transaction(store.url,write=True) as db:
            db.execute("UPDATE research_cost_records SET summary_payload='{}' WHERE sequence=%s",(sequence,))


def test_overview_projection_does_not_mutate_shared_history(departments):
    from api.research_history import linked_histories
    _,_,_,store,user,_=departments
    with transaction(store.url) as db:
        keys=[r[0] for r in db.execute('SELECT id FROM research_formulas')]
        histories=linked_histories(db,keys,summary=True)
    original=deepcopy(histories)
    store.listing(histories=histories)
    store.listing(user,histories=histories)
    assert histories==original
