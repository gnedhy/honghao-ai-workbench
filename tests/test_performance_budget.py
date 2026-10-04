"""Fixed synthetic load on the fixture-owned database; no production connection or business writes."""
from contextvars import ContextVar
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from statistics import median
from time import perf_counter

import psycopg
from openpyxl import load_workbook

from api.postgres import transaction
from api.procurement import ProcurementStore
from api.research import ResearchStore, capture, evaluate, packed
from api.sales_calculator import CalculatorStore, Evaluation
from tests.test_sales import sales


def test_critical_store_performance(sales, monkeypatch):
    policy = json.loads(Path('scripts/performance-budgets.json').read_text(encoding='utf-8'))
    fixture, limits = policy['fixture'], policy['limits']
    assert fixture['apiSamples'] == 9 and fixture['apiProducts'] == 200 and fixture['apiMaterials'] == 2000 and fixture['apiHistoryPerProduct'] == 20
    assert fixture['apiFormulaLines'] == 12 and fixture['apiCalculatorPanels'] == 12
    settings, _, sales_store, actor = sales
    with transaction(settings.database_url, write=True) as db:
        db.execute('DELETE FROM research_cost_records')
        db.execute('DELETE FROM research_formulas')
        formulas, costs, graph = [], [], []
        for index in range(fixture['apiProducts']):
            graph.append(dict(id=f'recipe:performance-{index}', name=f'P{index:05d}', owner='合成负责人', kind='recipe', revision=1, yield_='1',
                lines=[dict(kind='material', ref=f'M{(index * 12 + line) % fixture["apiMaterials"]:05d}', code=f'M{(index * 12 + line) % fixture["apiMaterials"]:05d}', quantity='1', source_row=None) for line in range(fixture['apiFormulaLines'])]))
            graph[-1]['yield'] = graph[-1].pop('yield_')
        base_prices = {f'M{i:05d}': {'latest_price': '24', 'inventory_price': '12'} for i in range(fixture['apiMaterials'])}
        base_inputs = {'package': {'recipes': graph, 'materials': []}, 'prices': base_prices}
        baseline = evaluate(base_inputs)
        for index in range(fixture['apiProducts']):
            formula = graph[index]
            formulas.append((formula['id'], index, packed(formula)))
            for version in range(1, fixture['apiHistoryPerProduct'] + 1):
                payload = dict(formula, latest_cost=str(4 + version), inventory_cost='12', status='ready', recorded_at=f'2026-10-{version:02d}T00:00:00Z', event_id=f'performance-{version}', change={'percent': None, 'reason': '首次'})
                if version == fixture['apiHistoryPerProduct']:
                    payload.update(formula=formula, graph=[formula], latest=baseline['latest'][formula['id']], inventory=baseline['inventory'][formula['id']],
                        calculations={basis: {formula['id']: values[formula['id']]} for basis, values in baseline.items()})
                costs.append((payload['event_id'], formula['id'], f'fixed-{version}', packed(payload)))
        db.cursor().executemany("INSERT INTO research_formulas(id,position,formula,lifecycle,revision) VALUES(%s,%s,%s,'active',1)", formulas)
        db.cursor().executemany('INSERT INTO research_cost_records(event_id,product_id,signature,payload) VALUES(%s,%s,%s,%s)', costs)
        db.cursor().executemany('INSERT INTO procurement_materials(id,code,name,unit,latest_price,updated_at) VALUES(%s,%s,%s,%s,%s,%s)',
            [(f'performance-{i}', f'M{i:05d}', f'合成原料 {i}', 'kg', '24', '2026-10-01T00:00:00Z') for i in range(fixture['apiMaterials'])])
        db.execute("INSERT INTO procurement_price_batches(id,version,published_by,published_at,item_count) VALUES('performance-batch',1,'synthetic','2026-10-01T00:00:00Z',%s)", (fixture['apiMaterials'],))
        db.cursor().executemany("INSERT INTO procurement_price_batch_items(batch_id,material_id,code,name,unit,latest_price) VALUES('performance-batch',%s,%s,%s,'kg','24')",
            [(f'performance-{i}', f'M{i:05d}', f'合成原料 {i}') for i in range(fixture['apiMaterials'])])
        db.cursor().executemany("INSERT INTO procurement_inventory(material_id,quantity,price,raw_price,source_sha256,sheet,source_row,imported_by,imported_at) VALUES(%s,'100','12','12','synthetic','fixture',%s,'synthetic','2026-10-01T00:00:00Z')",
            [(f'performance-{i}', i + 1) for i in range(fixture['apiMaterials'])])
        db.execute('INSERT INTO research_meta(id,payload) VALUES(1,%s) ON CONFLICT(id) DO UPDATE SET payload=EXCLUDED.payload', (packed({'materials': [{'code': f'M{i:05d}', 'unit': 'kg'} for i in range(fixture['apiMaterials'])]}),))
        frozen = capture(db)
        original_history = sha256(''.join(row[0] for row in db.execute('SELECT payload FROM research_cost_records ORDER BY sequence')).encode()).hexdigest()
    changed = deepcopy(frozen)
    for row in changed['prices'].values():
        row['latest_price'] = '25'

    def prepare_recalculation(index):
        # Each sample begins with the same 4,000 original cost rows, then really writes 200 new rows.
        with transaction(settings.database_url, write=True) as db:
            db.execute("DELETE FROM research_cost_records WHERE event_id LIKE 'performance-recompute-%'")
            db.execute("DELETE FROM research_events WHERE status!='done' OR id LIKE 'performance-recompute-%'")
            db.execute("INSERT INTO research_events(id,recorded_at,reason,inputs,status,attempts) VALUES(%s,'2026-10-21T00:00:00Z','合成性能样本',%s,'pending',0)", (f'performance-recompute-{index}', packed(changed)))
            assert db.execute('SELECT count(*) FROM research_cost_records').fetchone()[0] == fixture['apiProducts'] * fixture['apiHistoryPerProduct']

    counter = ContextVar('performance_sql_count', default=None)
    original = psycopg.Cursor.execute
    original_many = psycopg.Cursor.executemany
    def execute(cursor, *args, **kwargs):
        active = counter.get()
        if active is not None:
            active['count'] += 1
        return original(cursor, *args, **kwargs)
    def executemany(cursor, query, params_seq, *args, **kwargs):
        active = counter.get()
        if active is not None:
            params_seq = list(params_seq)
            active['count'] += len(params_seq)
        return original_many(cursor, query, params_seq, *args, **kwargs)
    monkeypatch.setattr(psycopg.Cursor, 'execute', execute)
    monkeypatch.setattr(psycopg.Cursor, 'executemany', executemany)
    evaluation = Evaluation.model_validate({'source': {'kind': 'product', 'product_id': 'research:recipe:performance-199', 'basis': 'latest'},
        'panels': [{'id': f'panel-{i}', 'name': f'对比 {i}', 'mode': 'domestic_direct', 'steps': [{'id': 'margin', 'label': '加价', 'operation': '+', 'value': '1'}]} for i in range(fixture['apiCalculatorPanels'])]})
    operations = {
        'procurementOverview': lambda: ProcurementStore(settings.database_url).overview(),
        'researchListing': lambda: ResearchStore(settings.database_url).listing(),
        'salesCatalog': lambda: sales_store.products(actor),
        'calculator': lambda: CalculatorStore(sales_store).evaluate(evaluation, actor),
        'researchExport': lambda: ResearchStore(settings.database_url).export_costs(),
        'researchRecalculation': lambda: ResearchStore(settings.database_url).process_events(),
    }
    metrics, samples = {}, {}
    for name, operation in operations.items():
        durations, statements = [], []
        for index in range(fixture['apiSamples']):
            if name == 'researchRecalculation':
                prepare_recalculation(index)
            count = {'count': 0}; token = counter.set(count); started = perf_counter()
            try:
                result = operation()
            finally:
                duration = (perf_counter() - started) * 1000
                counter.reset(token)
            # Check the real result in every sample, not only timing or an empty success response.
            if name == 'procurementOverview':
                assert len(result['materials']) == fixture['apiMaterials']
            elif name == 'researchListing':
                assert len(result['products']) == fixture['apiProducts'] and all(row['latest_cost'] == str(4 + fixture['apiHistoryPerProduct']) for row in result['products'])
            elif name == 'salesCatalog':
                assert len(result['products']) == fixture['apiProducts'] + fixture['apiMaterials']
                assert next(row for row in result['products'] if row['id'] == 'research:recipe:performance-199')['latest_cost'] == str(4 + fixture['apiHistoryPerProduct'])
            elif name == 'calculator':
                assert len(result['results']) == fixture['apiCalculatorPanels'] and all(row['error'] is None for row in result['results']) and result['source']['cost'] == str(4 + fixture['apiHistoryPerProduct'])
            elif name == 'researchExport':
                book = load_workbook(BytesIO(result), read_only=True, data_only=False)
                try:
                    assert book.sheetnames == ['总览', '合成负责人']
                    assert book['总览'].max_row == fixture['apiProducts'] + 3
                    assert book['总览']['A203'].value == 'P00199' and book['总览']['C4'].value == 24 and book['总览']['D4'].value == 12
                    assert book['合成负责人'].max_row >= fixture['apiProducts'] * fixture['apiFormulaLines']
                finally:
                    book.close()
            else:
                assert result == 1
                with transaction(settings.database_url) as db:
                    rows = db.execute('SELECT payload FROM research_cost_records WHERE event_id=%s', (f'performance-recompute-{index}',)).fetchall()
                    assert len(rows) == fixture['apiProducts']
                    assert all(json.loads(row[0])['latest_cost'] == '25' and len(json.loads(row[0])['latest']['lines']) == fixture['apiFormulaLines'] for row in rows)
                    assert db.execute('SELECT status FROM research_events WHERE id=%s', (f'performance-recompute-{index}',)).fetchone()[0] == 'done'
                    preserved = ''.join(row[0] for row in db.execute("SELECT payload FROM research_cost_records WHERE event_id NOT LIKE 'performance-recompute-%' ORDER BY sequence"))
                    assert sha256(preserved.encode()).hexdigest() == original_history
            durations.append(duration); statements.append(count['count'])
        assert len(durations) == fixture['apiSamples'] and min(statements) > 0
        metrics.update({f'api.{name}.medianMs': median(durations), f'api.{name}.maxMs': max(durations), f'api.{name}.sqlCalls': max(statements)})
        samples[name] = {'milliseconds': durations, 'sqlCalls': statements}
    output = Path('.scratch/performance-budget/api.json'); output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps({'schema': 1, 'kind': 'api', 'fixture': fixture, 'metrics': metrics, 'samples': samples}, indent=2) + '\n', encoding='utf-8')
    assert set(metrics) == {key for key in limits if key.startswith('api.')}, 'Incomplete API performance measurements'
    for key, value in metrics.items():
        assert value <= limits[key], f'{key}: {value:.2f} exceeds {limits[key]}'
