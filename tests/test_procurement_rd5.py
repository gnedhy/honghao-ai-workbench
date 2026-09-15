import copy
import json
import os
import psycopg
from api.postgres import migrate, transaction
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from api.cli import main
from api.identity import IdentityStore
from api.procurement import ProcurementStore
from api.procurement_excel import read_workbook
from api.procurement_rd5 import export_preparation, import_rd5, preview, revise_preparation
from api.research_formulas import DEFAULT_COMPOSITE, calculate, parse_workbook
from tests.helpers import authenticated_client
from tests.test_procurement_isolation import import_prices, publish
from tests.test_procurement_workbench import procurement_settings


def make_source(path):
    w = Workbook()
    s = w.active
    s.title = "原料单价"
    s.append(["编号", "核算价格", "最新价格", "库存价格", "在途价格"])
    for r, (code, latest, stock) in enumerate([
        ("A", 5, 4), ("B", None, 8), ("CF026", 12.7, 11.5), ("CF026K", 12.9, 14.5),
        ("CF020C", None, None), ("CF020D", 9.1, 8.9), ("CF401B", 12.62, 14.06),
        ("纯水", 0, 0), ("CF011", None, None), ("CF007B", "=C2", "=D2"),
        ("CF063A", "=C2/2", "=D2/2"),
    ], 2):
        s.cell(r, 1, code)
        if latest is not None or stock is not None:
            s.cell(r, 2, f'=IF(ISBLANK(C{r}),IF(ISBLANK(E{r}),IF(ISBLANK(D{r}),"check",D{r}),E{r}),IF(ISNUMBER(C{r}),C{r},"check"))')
        s.cell(r, 3, latest)
        s.cell(r, 4, stock)
    for r, codes in zip([20, 24, 28, 32], [("CF026", "CF020C"), ("CF026", "CF020D"), ("CF026K", "CF020C"), ("CF026K", "CF020D")]):
        s.cell(r, 8, "CF401B")
        for x, code, amount in zip([r, r + 1], codes, [3000, 300]):
            s.cell(x, 9, code)
            s.cell(x, 10, f"=L{x}/SUM(L{r}:L{r+1})")
            s.cell(x, 11, f"=VLOOKUP(I{x},原料单价!A:B,2,0)")
            s.cell(x, 12, amount)
        s.cell(r, 13, 0.995)
        s.cell(r, 14, f"=SUMPRODUCT(J{r}:J{r+1},K{r}:K{r+1})/M{r}")
    s = w.create_sheet("测试负责（配方及成本）")
    s.append(["配方名称", "原料名称", "配方比例%", "原料单价（元/kg）", "实际投料（kg）", "收率%", "成本（元/kg）"])
    for start, name, items, yield_ in [
        (2, "K172-C", [("CF401B", 10), ("B", 10), ("A", 10)], .9),
        (6, "RH-1A", [("K172-C", 20), ("纯水", 40), ("纯水", 40)], .8),
        (10, "缺价方案", [("CF401B", 10)], .99),
    ]:
        end = start + len(items) - 1
        for r, (code, amount) in enumerate(items, start):
            formula = f"=VLOOKUP(B{r},A:G,7,0)" if code == "K172-C" else "=原料单价!N20" if name == "缺价方案" else f"=VLOOKUP(B{r},原料单价!A:B,2,0)"
            for col, value in enumerate([name, code, f"=E{r}/SUM(E${start}:E${end})", formula, amount], 1):
                s.cell(r, col, value)
        s.cell(start, 6, yield_)
        s.cell(start, 7, f"=SUMPRODUCT(C{start}:C{end},D{start}:D{end})/F{start}")
    w.save(path)


@pytest.fixture
def data(tmp_path):
    settings = procurement_settings(tmp_path)
    with authenticated_client(settings) as client:
        assert publish(client, import_prices(client, "2026-09-08", [("A", 30)])).status_code == 200
        admin = client.get("/api/me").json()["id"]
        with transaction(settings.database_url, write=True) as db:
            a_id = db.execute("SELECT id FROM procurement_materials WHERE code='A'").fetchone()[0]
            db.execute("INSERT INTO procurement_departments (id, name, position) VALUES ('rd5','研发五部',4)")
            db.execute("INSERT INTO procurement_department_materials (department_id, material_id) VALUES ('rd5',%s)", (a_id,))
            db.execute("INSERT INTO procurement_inventory (material_id, quantity, price, raw_price, source_sha256, sheet, source_row, imported_by, imported_at) VALUES (%s, '12.345', '9', '9', 'original', '原料行情总表', 2, %s, 'original')", (a_id, admin))
            db.execute("INSERT INTO procurement_materials(id,code,name,unit,updated_at,archived_at) VALUES ('old-b','B','B','kg','original','archived')")
        source = tmp_path / 'rd5.xlsx'
        make_source(source)
        yield settings, client, source, admin


def apply(data, plan=None, actor=None):
    settings, _, source, admin = data
    plan = plan or preview(settings.database_url, source.read_bytes())
    return import_rd5(settings.database_url, source, actor or admin, expected_sha256=plan["sha256"],
                      expected_state_sha256=plan["state_sha256"], effective_date=date(2026, 9, 11), data_dir=settings.data_dir)


def table_rows(path):
    with transaction(path) as db:
        return {name: db.execute(f'SELECT * FROM {name} ORDER BY _order').fetchall() for (name,) in
                db.execute("SELECT table_name FROM information_schema.tables WHERE table_schema='public' AND table_name LIKE 'procurement_%'")}


@pytest.mark.parametrize('updated_price, expected_change', [(30, 0), (33, 0.1)])
def test_supplement_preserves_old_comparisons_and_next_price_update_uses_current_price(data, updated_price, expected_change):
    settings, client, source, _ = data
    assert publish(client, import_prices(client, '2026-09-10', [('A', updated_price)])).status_code == 200
    before = client.get('/api/workbenches/procurement/overview').json()
    old = next(m for m in before['materials'] if m['code'] == 'A')
    mid = old['id']
    original = before['ledger_comparison']['items'][mid]
    assert original['change'] == expected_change
    plan = preview(settings.database_url, source.read_bytes())
    apply(data, plan)
    after = client.get('/api/workbenches/procurement/overview').json()
    retained = next(m for m in after['materials'] if m['id'] == mid)
    assert after['ledger_comparison']['items'][mid] == original
    for key in ['published_price', 'previous_published_price', 'published_price_date', 'price_modifier', 'inventory_price', 'inventory_quantity']:
        assert retained[key] == old[key], key
    assert after['batches'][0]['comparison']['items'][mid]['change'] == 0
    detail = client.get(f'/api/workbenches/procurement/materials/{mid}').json()
    assert detail['comparison'] == original
    new_id = next(m['id'] for m in after['materials'] if m['code'] == 'CF026')
    assert after['ledger_comparison']['items'][new_id]['kind'] == 'first'
    assert after['ledger_comparison']['items'][new_id]['change'] is None
    additions = after['batches'][0]['comparison']['added_material_ids']
    assert mid not in additions and new_id in additions and len(additions) == 7
    version = client.get('/api/workbenches/procurement/batches/' + after['batches'][0]['id']).json()
    assert version['comparison']['added_material_ids'] == additions
    assert len(version['items']) == 8  # The immutable snapshot still carries the existing material.

    assert apply(data, plan)['already_imported'] is True
    assert ProcurementStore(settings.database_url).batch_comparison(before['batches'][0]['id']) == before['batches'][0]['comparison']

    # A later addition through the ordinary importer preserves both old and newly added baselines.
    assert publish(client, import_prices(client, '2026-09-12', [('EXTRA', 8)])).status_code == 200
    extra = client.get('/api/workbenches/procurement/overview').json()
    assert extra['ledger_comparison']['items'][mid] == original
    assert extra['ledger_comparison']['items'][new_id] == after['ledger_comparison']['items'][new_id]

    # Ordinary price updates still compare with the prior formal price, including unchanged quotes.
    assert publish(client, import_prices(client, '2026-09-13', [('A', updated_price)])).status_code == 200
    latest = client.get('/api/workbenches/procurement/overview').json()
    entry = latest['ledger_comparison']['items'][mid]
    assert entry['change'] == 0
    assert entry['previous'] == str(updated_price)
    assert entry['previous_version'] == extra['batches'][0]['version']


def test_import_preserves_existing_data_restores_ids_publishes_and_calculates(data, tmp_path):
    settings, client, source, admin = data
    plan = preview(settings.database_url, source.read_bytes())
    assert plan['counts'] == {'keep': 1, 'restore': 1, 'create': 8, 'omit': 1, 'active_before': 1, 'active_after': 10,
                              'department_before': 1, 'department_after': 10, 'latest_prices': 7, 'inventory_prices': 7,
                              'recipes': 3, 'recipe_lines': 7, 'composites': 4}
    before = table_rows(settings.database_url)
    result = apply(data, plan)
    assert result['batch']['version'] == 2
    assert result['blocked_recipes'] == ['缺价方案']
    after = table_rows(settings.database_url)
    assert after['procurement_materials'][0] == before['procurement_materials'][0]
    assert after['procurement_inventory'][0] == before['procurement_inventory'][0]
    assert after['procurement_price_batches'][:1] == before['procurement_price_batches']
    assert after['procurement_price_batch_items'][:1] == before['procurement_price_batch_items']
    overview = client.get('/api/workbenches/procurement/overview').json()
    rows = {m['code']: m for m in overview['materials']}
    assert rows['A']['published_price'] == '30'
    assert rows['B']['id'] == 'old-b' and rows['B']['inventory_price'] == '8'
    assert rows['B']['published_price'] is None and rows['B']['inventory_quantity'] is None
    assert rows['纯水']['published_price'] == '0' and rows['纯水']['inventory_price'] is None
    assert rows['CF007B']['published_price'] == '5'
    assert rows['CF063A']['published_price'] == '2.5'
    assert all(m['published_price_date'] == '2026-09-11' for k, m in rows.items() if k != 'A' and m['published_price'] is not None)
    assert all(m['price_modifier']['id'] == admin for k, m in rows.items() if k != 'A' and m['published_price'] is not None)
    export = export_preparation(settings.database_url, result['sha256'], tmp_path / 'review')
    with open(export['package'], encoding='utf-8') as stream:
        package = json.load(stream)
    calc = package['current_calculation']
    composite = (Decimal(3000) * Decimal('12.9') + Decimal(300) * Decimal('9.1')) / (Decimal(3300) * Decimal('.995'))
    expected = (composite + Decimal(8) + Decimal(30)) / 3 / Decimal('.9') / 4
    assert abs(Decimal(calc['recipe:RH-1A']['cost']) - expected) < Decimal('1e-20')
    assert calc['recipe:缺价方案']['missing_materials'] == ['CF020C']
    assert calc['recipe:K172-C']['lines'][0]['ref'] == DEFAULT_COMPOSITE
    assert package['source_replay']['recipe:缺价方案']['cost'] is not None
    assert 'source_blank_as_zero' in json.dumps(package['source_replay'])
    assert package['recipes'][0]['lines'][0]['material_id'] == rows['CF401B']['id']
    assert apply(data, plan)['already_imported'] is True
    assert table_rows(settings.database_url) == after
    # Later normal price activation must leave independent inventory intact.
    newer = import_prices(client, '2026-09-12', [('B', 8)])
    assert publish(client, newer).status_code == 200
    rows = {m['code']: m for m in client.get('/api/workbenches/procurement/overview').json()['materials']}
    assert (rows['B']['inventory_price'], rows['B']['inventory_quantity']) == ('8', None)


@pytest.mark.parametrize('target,value', [('A3','A'), ('C4',-2), ('C4','=HYPERLINK("https://example.com")'), ('B4','=MIN(C4,D4)'), ('E4',4)])
def test_bad_source_rejected_without_writes(data, target, value):
    settings, _, source, _ = data
    before = table_rows(settings.database_url)
    w = load_workbook(source)
    w['原料单价'][target] = value
    w.save(source)
    with pytest.raises(ValueError):
        apply(data)
    assert table_rows(settings.database_url) == before


@pytest.mark.parametrize('target,value', [('E2',-1), ('F2',0), ('D2','=VLOOKUP(B2,其他!A:B,2,0)'), ('C2',.5)])
def test_invalid_recipe_rejected(data, target, value):
    _, _, source, _ = data
    w = load_workbook(source)
    w['测试负责（配方及成本）'][target] = value
    w.save(source)
    with pytest.raises(ValueError):
        parse_workbook(source.read_bytes())


@pytest.mark.parametrize('mutation', ['material', 'prices', 'department', 'alias'])
def test_stale_preflight_stops_before_writing(data, mutation):
    settings, _, source, _ = data
    plan = preview(settings.database_url, source.read_bytes())
    with transaction(settings.database_url, write=True) as db:
        db.execute({'material': "UPDATE procurement_materials SET code='A-new' WHERE code='A'",
                    'prices': "UPDATE procurement_price_batch_items SET latest_price='31'",
                    'department': "DELETE FROM procurement_department_materials",
                    'alias': "INSERT INTO procurement_material_aliases (alias_code, material_id, created_by, created_at) VALUES ('CF026','old-b','someone','now')"}[mutation])
    before = table_rows(settings.database_url)
    with pytest.raises(ValueError):
        apply(data, plan)
    assert table_rows(settings.database_url) == before


def test_file_change_and_pending_round_stop_import(data):
    settings, client, source, _ = data
    plan = preview(settings.database_url, source.read_bytes())
    source.write_bytes(source.read_bytes() + b'changed')
    with pytest.raises(ValueError, match='文件已变化'):
        apply(data, plan)
    make_source(source)
    import_prices(client, '2026-09-09', [('A', 31)])
    with pytest.raises(ValueError, match='未完成'):
        apply(data)


def test_entire_transaction_and_archive_roll_back(data):
    settings, _, _, _ = data
    with transaction(os.environ['HONGHAO_TEST_MIGRATION_URL'], write=True) as db:
        db.execute("""CREATE FUNCTION fail_receipt() RETURNS trigger LANGUAGE plpgsql AS $$
            BEGIN RAISE EXCEPTION 'late failure' USING ERRCODE='23514'; END $$;
            CREATE TRIGGER fail_receipt BEFORE INSERT ON procurement_rd5_imports
            FOR EACH ROW EXECUTE FUNCTION fail_receipt()""")
    before = table_rows(settings.database_url)
    with pytest.raises(psycopg.IntegrityError, match='late failure'):
        apply(data)
    assert table_rows(settings.database_url) == before
    assert not list((settings.data_dir / 'controlled-work' / 'procurement-sources').glob('*.xlsx'))


def test_permissions_cli_and_generic_formula_rejection(data, capsys):
    settings, _, source, _ = data
    user = IdentityStore(settings.database_url).create_user(username='reader', display_name='reader', department=None,
                                                           password='Reader-password-2026', scope_levels={'procurement': 4})
    with pytest.raises(PermissionError):
        apply(data, actor=user['id'])
    assert main(['procurement-rd5', str(source)], settings=settings) == 0
    assert json.loads(capsys.readouterr().out)['counts']['restore'] == 1
    assert main(['procurement-rd5', str(source), '--apply'], settings=settings) == 1
    assert '预检' in capsys.readouterr().err
    with pytest.raises(ValueError, match='公式'):
        read_workbook(source.read_bytes())


def test_quantity_nullable_schema_preserves_source_and_preferences(data):
    settings, _, _, admin = data
    before = table_rows(settings.database_url)['procurement_inventory']
    preferences = ProcurementStore(settings.database_url).preferences(admin)
    assert migrate(os.environ['HONGHAO_TEST_MIGRATION_URL'], 'test') == 2
    assert migrate(os.environ['HONGHAO_TEST_MIGRATION_URL'], 'test') == 2
    assert ProcurementStore(settings.database_url).schema_version() == 7
    assert table_rows(settings.database_url)['procurement_inventory'] == before
    assert ProcurementStore(settings.database_url).preferences(admin) == preferences
    with transaction(settings.database_url) as db:
        assert db.execute("SELECT is_nullable FROM information_schema.columns WHERE table_schema='public' AND table_name='procurement_inventory' AND column_name='quantity'").fetchone() == ('YES',)


def test_graph_cycles_missing_propagation_and_zero_quantities(tmp_path):
    source = tmp_path / 'source.xlsx'
    make_source(source)
    p = parse_workbook(source.read_bytes())
    prices = {m['code']: m for m in p['materials']}
    costs = calculate(p, prices)
    assert costs['recipe:RH-1A']['cost'] is not None
    prices.pop('B')
    assert calculate(p, prices)['recipe:RH-1A']['missing_materials'] == ['B']
    prices['B'] = {'latest_price': '100', 'inventory_price': '1', 'in_transit_price': '0.1'}
    assert calculate(p, prices)['recipe:K172-C']['lines'][1]['unit_cost'] == '100'
    p = copy.deepcopy(p)
    p['recipes'][0]['lines'][0].update(kind='recipe', ref='recipe:RH-1A')
    with pytest.raises(ValueError, match='循环'):
        calculate(p, prices)


@pytest.fixture
def trial_data(data):
    settings, _, source, _ = data
    w = load_workbook(source)
    s = w['测试负责（配方及成本）']
    for r, variant, cost_row in [(10, '026+020C', 20), (14, '026+020D', 24), (18, '026K+020C', 28)]:
        for col, value in enumerate([f'K172-C（{variant}）', 'CF401B', f'=E{r}/SUM(E${r}:E${r})',
                                    f'=原料单价!N{cost_row}', 10, .99, f'=SUMPRODUCT(C{r}:C{r},D{r}:D{r})/F{r}'], 1):
            s.cell(r, col, value)
    w.save(source)
    result = apply(data)
    return data, result['sha256']


def test_current_scope_preserves_procurement_source_and_reimport(trial_data, tmp_path, capsys):
    data, digest = trial_data
    settings, _, source, admin = data
    before = table_rows(settings.database_url)
    original = json.loads(before['procurement_rd5_imports'][0][4])
    plan = revise_preparation(settings.database_url, digest)
    assert plan['counts'] == {'recipes': 2, 'recipe_lines': 6, 'composites': 1,
                              'historical_recipes': 3, 'historical_composites': 3}
    assert plan['blocked_recipes'] == []
    assert table_rows(settings.database_url) == before
    assert main(['procurement-rd5-current', digest, '--apply'], settings=settings) == 1
    assert '预检' in capsys.readouterr().err
    assert main(['procurement-rd5-current', digest, '--apply', '--actor-id', admin,
                 '--state-sha256', plan['state_sha256'], '--export-dir', str(tmp_path / 'revised')], settings=settings) == 0
    exported = json.loads(capsys.readouterr().out)['export']
    package = json.loads((tmp_path / 'revised' / 'rd5-preparation.json').read_text(encoding='utf-8'))
    assert exported['current_counts'] == {'recipes': 2, 'recipe_lines': 6, 'composites': 1}
    assert package['source_replay'] == original['source_replay']
    assert package['recipes'] == original['recipes']
    assert package['materials'] == original['materials']
    assert calculate(package, {}, source_mode=True) == original['source_replay']
    assert len(package['current_calculation']) == 3
    for key, current in package['current_calculation'].items():
        assert current == original['current_calculation'][key]
    after = table_rows(settings.database_url)
    for name in before.keys() - {'procurement_rd5_imports', 'procurement_admin_events'}:
        assert before[name] == after[name], name
    assert before['procurement_rd5_imports'][0][:4] == after['procurement_rd5_imports'][0][:4]
    assert len(after['procurement_admin_events']) == len(before['procurement_admin_events']) + 1
    plan = revise_preparation(settings.database_url, digest)
    assert plan['already_current'] is True
    revise_preparation(settings.database_url, digest, actor_id=admin, expected_state_sha256=plan['state_sha256'])
    assert import_rd5(settings.database_url, source, admin, expected_sha256=digest,
                      expected_state_sha256='unused', effective_date=date(2026, 9, 11), data_dir=settings.data_dir)['already_imported']
    assert table_rows(settings.database_url) == after
    review = (tmp_path / 'revised' / '研发五部数据核对.md').read_text(encoding='utf-8')
    assert '历史试算（不列入当前有效配方）' in review and 'CF020C 已停购' in review
    assert '两种相关K172-C配方及两种复配方案待补价' not in review


@pytest.mark.parametrize('failure', ['permission', 'state', 'package', 'source', 'rollback'])
def test_current_revision_guards_and_rollback(trial_data, failure):
    data, digest = trial_data
    settings, _, _, admin = data
    plan = revise_preparation(settings.database_url, digest)
    if failure == 'permission':
        admin = 'not-an-admin'
    elif failure in {'state', 'package'}:
        with transaction(settings.database_url, write=True) as db:
            db.execute({'state': "UPDATE procurement_inventory SET price='10'",
                        'package': "UPDATE procurement_rd5_imports SET package_json=package_json || ' '"}[failure])
    elif failure == 'rollback':
        with transaction(os.environ['HONGHAO_TEST_MIGRATION_URL'], write=True) as db:
            db.execute("""CREATE FUNCTION fail_revision() RETURNS trigger LANGUAGE plpgsql AS $$
                BEGIN RAISE EXCEPTION 'late failure' USING ERRCODE='23514'; END $$;
                CREATE TRIGGER fail_revision BEFORE INSERT ON procurement_admin_events
                FOR EACH ROW EXECUTE FUNCTION fail_revision()""")
    elif failure == 'source':
        (settings.data_dir / 'controlled-work' / 'procurement-sources' / f'{digest}.xlsx').write_bytes(b'changed')
    before = table_rows(settings.database_url)
    with pytest.raises((PermissionError, ValueError, psycopg.IntegrityError)):
        revise_preparation(settings.database_url, digest, actor_id=admin, expected_state_sha256=plan['state_sha256'])
    assert table_rows(settings.database_url) == before


def test_replacement_uses_current_d_price_and_propagates_missing(trial_data):
    data, digest = trial_data
    settings, _, _, admin = data
    plan = revise_preparation(settings.database_url, digest)
    revise_preparation(settings.database_url, digest, actor_id=admin, expected_state_sha256=plan['state_sha256'])
    package = json.loads(table_rows(settings.database_url)['procurement_rd5_imports'][0][4])
    # A source CF020C line remains traceable; current calculation resolves its replacement once.
    line = package['recipes'][0]['lines'][1]
    line.update(code='CF020C', ref='CF020C')
    prices = package['price_inputs']
    prices['CF020C'] = {'latest_price': '1', 'inventory_price': '2'}
    prices['CF020D'] = {'latest_price': '9.3', 'inventory_price': '8.9'}
    calc = calculate(package, prices)
    current = calc['recipe:K172-C']['lines'][1]
    assert (current['code'], current['ref'], current['unit_cost'], current['basis']) == ('CF020C', 'CF020D', '9.3', 'latest')
    prices['CF020D']['latest_price'] = None
    assert calculate(package, prices)['recipe:K172-C']['lines'][1]['unit_cost'] == '8.9'
    prices.pop('CF020D')
    assert calculate(package, prices)['recipe:RH-1A']['missing_materials'] == ['CF020D']
    assert calculate(package, prices)['recipe:RH-1A']['cost'] is None
    line.update(kind='recipe', ref='recipe:K172-C（026+020C）')
    with pytest.raises(ValueError, match='历史试算'):
        calculate(package, prices)
