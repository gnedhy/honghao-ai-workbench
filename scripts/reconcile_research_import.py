"""Refresh independent evidence and trial costs; never writes application data."""
from __future__ import annotations

import argparse
import csv
import difflib
import hashlib
import html
import json
import shutil
import sqlite3
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import date, datetime, timezone, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from analyze_research_import import normalized_feed, packed, render_evidence, RULES
from api.research_formulas import calculate, ERP_CONCENTRATIONS, decimal_text


WAREHOUSES = {
    '宏昊原料仓', '宏昊成品仓', '浙江原料仓', '浙江成品仓',
    '宏图原料仓', '宏图成品仓', '宏图成品尾数仓', '宏图车间仓',
}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def formula_digest(formula):
    return hashlib.sha256(json.dumps(formula, ensure_ascii=False, sort_keys=True).encode('utf-8')).hexdigest()


def write_csv(path, headers, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(rows)


def signature(lines):
    # Record order remains meaningful; drop only inert, zero-ratio rows in equivalence checks.
    return tuple((l['normalized_feed'], Decimal(l['quantity']), Decimal(l['ratio']))
                 for l in lines if not l['excluded'] and
                 (Decimal(l['quantity']) != 0 or Decimal(l['ratio']) != 0))


def summarize_inputs(formula):
    amounts = defaultdict(Decimal)
    for line in formula['lines']:
        amounts[normalized_feed(line['code'])] += Decimal(line['quantity'])
    return {k: v for k, v in amounts.items() if v != 0}


def price_codes(line):
    refs = line.get('price_refs') or [line.get('price_ref', line['ref'])]
    return {'CF232' if ref in ERP_CONCENTRATIONS else ref for ref in refs}


def reconcile(root: Path, prior: Path, output: Path, live_path: Path):
    if output.exists():
        raise ValueError('使用新的输出目录，保留此前分析证据')
    decisions_path = prior / '业务确认-20261007.json'
    decisions = json.loads(decisions_path.read_text(encoding='utf-8'))
    if any(decisions.get(key) for key in ('product_import_exclusions', 'product_code_corrections', 'recipe_input_selections', 'inventory_price_reviews')):
        raise ValueError('已有后续范围及取价确认；此原始梳理入口不能覆盖已确认候选。使用当前独立候选和finalize_research_import.py核验后续调整。')
    stock_inputs = {(r['product'], r['source_row']): r for r in decisions.get('stock_input_selections', [])}
    assert len(stock_inputs) == len(decisions.get('stock_input_selections', []))
    applied_stock_inputs = set()
    replacements = decisions.get('material_substitutions', {})
    applied_replacements = Counter()
    zero_inputs = decisions.get('unpriced_zero_cost_inputs', {})
    applied_zero_inputs = Counter()
    yield_selections = decisions.get('yield_selections', {})
    applied_yield_selections = set()
    system_confirmation = decisions.get('system_formula_replacement_confirmation', {})
    system_selections = system_confirmation.get('products', {})
    applied_system_selections = set()
    erp_adoption_formulas = {}
    inventory_confirmation = decisions.get('inventory_confirmation', {})
    shared_prices = inventory_confirmation.get('price_scope') == 'same_erp_item_across_companies_factories'
    cutoff = inventory_confirmation.get('quantity_asof')
    if cutoff:
        assert date.fromisoformat(cutoff).isoformat() == cutoff
    live = json.loads(live_path.read_text(encoding='utf-8'))
    assert live['read_only'] and live['database'] == 'honghao_production'
    output.mkdir(parents=True)
    originals = output / '来源原件'
    originals.mkdir()
    old = sqlite3.connect('file:' + str(prior / '研发数据分析.db') + '?mode=ro', uri=True)
    db = sqlite3.connect(output / '研发数据分析.db')
    old.backup(db)
    old.close()
    db.row_factory = sqlite3.Row
    for source in list(db.execute('select * from sources')):
        directory = '研发产品相关' if source['kind'] in {'bom', 'owners', 'yields'} else '采购原料相关'
        path = root / directory / source['filename']
        assert digest(path) == source['sha256'], f"原来源变化：{source['filename']}"
        shutil.copy2(path, originals / path.name)
    for item in decisions['sources']:
        assert db.execute('select sha256 from sources where kind=?', (item['kind'],)).fetchone()[0] == item['sha256']

    for batch in decisions.get('owner_confirmation_batches', []):
        path = prior / '来源确认' / batch['filename']
        assert digest(path) == batch['sha256']
        wb = load_workbook(path, read_only=True, data_only=True)
        raw = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = wb[batch['sheet']]
            values, formulas = list(sheet.values), list(raw[sheet.title].values)
            rows = [(n, list(r)) for n, r in enumerate(values, 1) if r and isinstance(r[0], (int, float))]
            assert len(rows) == batch['rows']
            kind = 'owner_confirmation'
            db.execute('insert into sources values(?,?,?,?,?,?)', (kind, path.name, batch['sha256'], sheet.title, packed(values[4]), len(rows)))
            db.executemany('insert into raw_rows values(?,?,?,?)', [(kind, n, packed(r), packed(formulas[n - 1])) for n, r in rows])
        finally:
            wb.close()
            raw.close()
        shutil.copy2(path, originals / path.name)

    for batch in decisions.get('material_price_confirmation_batches', []):
        path = prior / '来源确认' / batch['filename']
        assert digest(path) == batch['sha256']
        wb = load_workbook(path, read_only=True, data_only=True)
        raw = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = wb[batch['sheet']]
            values, formulas = list(sheet.values), list(raw[sheet.title].values)
            rows = [(n, list(r)) for n, r in enumerate(values, 1) if r and isinstance(r[0], (int, float))]
            assert len(rows) == batch['rows']
            kind = 'material_price_confirmation_' + batch['sha256'][:12]
            db.execute('insert into sources values(?,?,?,?,?,?)', (kind, path.name, batch['sha256'], sheet.title, packed(values[4]), len(rows)))
            db.executemany('insert into raw_rows values(?,?,?,?)', [(kind, n, packed(r), packed(formulas[n - 1])) for n, r in rows])
        finally:
            wb.close()
            raw.close()
        shutil.copy2(path, originals / path.name)

    current = {}
    files = {
        'erp_prices': '1007 原料库存单价明细表（信息补充）.xlsx',
        'erp_stock': '1007 原料库存明细表(带仓库信息).xlsx',
        'erp_prices_compact': '1007 原料库存单价明细表.xlsx',
        'erp_stock_compact': '1007 原料库存明细表.xlsx',
    }
    for kind, filename in files.items():
        path = root / '采购原料相关' / filename
        wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        raw = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            assert len(wb.sheetnames) == 1
            sheet = wb.active
            values = list(sheet.values)
            formulas = list(raw[sheet.title].values)
            rows = [(n, list(r)) for n, r in enumerate(values[1:], 2) if any(v is not None for v in r)]
            db.execute('insert into sources values(?,?,?,?,?,?)',
                       (kind, filename, digest(path), sheet.title, packed(values[0]), len(rows)))
            db.executemany('insert into raw_rows values(?,?,?,?)',
                           [(kind, n, packed(r), packed(formulas[n - 1])) for n, r in rows])
            current[kind] = rows
        finally:
            wb.close()
            raw.close()
        shutil.copy2(path, originals / filename)
    assert [(r[9], r[12]) for _, r in current['erp_prices']] == [tuple(r) for _, r in current['erp_prices_compact']]
    assert [(r[5], r[7], r[9]) for _, r in current['erp_stock']] == [tuple(r) for _, r in current['erp_stock_compact']]
    assert len({str(r[8]) for _, r in current['erp_prices']}) == len(current['erp_prices'])
    assert len({tuple(str(r[k]) for k in (0, 2, 4, 13)) for _, r in current['erp_stock']}) == len(current['erp_stock'])

    db.executescript('''
        ALTER TABLE review_issues RENAME TO initial_review_issues;
        CREATE TABLE review_issues(issue_id INTEGER PRIMARY KEY,category TEXT,product TEXT,evidence_json TEXT,status TEXT DEFAULT '待确认',decision TEXT DEFAULT '');
        CREATE TABLE business_decisions(payload_json TEXT);
        CREATE TABLE resolved_formulas(product TEXT PRIMARY KEY,formula_json TEXT,source_blocks_json TEXT,selection_basis TEXT);
        CREATE TABLE inventory_rows(source_row INTEGER PRIMARY KEY,code TEXT,item TEXT,company TEXT,factory TEXT,warehouse TEXT,quantity TEXT,unit TEXT,price TEXT,price_source_row INTEGER,price_company TEXT,price_factory TEXT,cross_entity INTEGER);
        CREATE TABLE material_candidates(code TEXT PRIMARY KEY,unit TEXT,quantity TEXT,erp_inventory_price TEXT,latest_price TEXT,evidence_json TEXT);
        CREATE TABLE trial_costs(product TEXT,policy TEXT,cost TEXT,missing_json TEXT,status TEXT,result_json TEXT,PRIMARY KEY(product,policy));
        CREATE TABLE system_differences(product TEXT PRIMARY KEY,classification TEXT,evidence_json TEXT);
    ''')
    db.execute('insert into business_decisions values(?)', (packed(decisions),))
    db.execute('insert into current_snapshot values(?,?)', ('20261007_live_read_only', packed(live)))
    rules = dict(RULES, inventory_status='1007来源已分析；剩余业务疑点确认后统一写入', inventory_scope='仅当前BOM引用的化工投料，生产相关仓库', erp_source_label='ERP 库存', inventory_quantity_asof=cutoff, inventory_price_scope='同一ERP品号跨公司和工厂统一取价' if shared_prices else '待确认')
    db.execute('delete from confirmed_rules')
    db.execute('insert into confirmed_rules values(?)', (packed(rules),))
    issues = []

    def issue(category, product, evidence):
        issues.append((category, product, evidence))
        db.execute('insert into review_issues(category,product,evidence_json) values(?,?,?)',
                   (category, product, packed(evidence)))

    groups = defaultdict(list)
    for block in db.execute('select * from bom_blocks order by start_row'):
        groups[block['product']].append(dict(block, lines=[dict(r) for r in db.execute(
            'select * from bom_lines where block_id=? order by source_row', (block['block_id'],))]))
    owners = defaultdict(list)
    for row in db.execute('select * from product_owners order by source_row'):
        owners[row['product']].append(dict(row))
    yields = {r['product']: dict(r) for r in db.execute('select * from yield_summary')}
    ids = {r['name']: r['id'] for r in live['formulas']}
    ids.update({code: ids.get(code, 'recipe:' + code) for code in groups})
    hashes = dict(db.execute('select kind,sha256 from sources'))
    resolved = {}
    refs = set()
    for code, blocks in groups.items():
        selection = decisions['formula_selections'].get(code, {})
        if 'bom_header_row' in selection:
            choices = [b for b in blocks if b['start_row'] == selection['bom_header_row']]
            assert len(choices) == 1
            block, basis = choices[0], '用户指定版本'
        elif selection.get('ratio_policy') == 'quantity_share':
            assert {b['start_row'] for b in blocks} == set(selection['source_bom_header_rows'])
            signatures = {tuple((l['normalized_feed'], Decimal(l['quantity'])) for l in b['lines'] if not l['excluded']) for b in blocks}
            assert len(signatures) == 1
            block, basis = deepcopy(blocks[0]), '用户确认重算比例；原始版本全部保留'
            total = sum(Decimal(l['quantity']) for l in block['lines'] if not l['excluded'])
            assert total == Decimal(selection['ratio_denominator_quantity'])
            for line in block['lines']:
                if not line['excluded']:
                    line['ratio'] = str(Decimal(line['quantity']) / total)
        else:
            signatures = {signature(b['lines']) for b in blocks}
            if len(signatures) != 1:
                issue('未确认配方版本', code, blocks)
                continue
            block, basis = blocks[0], '单一配方' if len(blocks) == 1 else '等价明细用于分析；原始版本全部保留'
        assigned = {(r['owner'], r['department']) for r in owners[code]}
        selected_owner = decisions['owner_selections'].get(code)
        if selected_owner:
            assignment = selected_owner['owner'], selected_owner['department']
            assert assignment in assigned
        elif len(assigned) == 1:
            assignment = next(iter(assigned))
        else:
            assignment = None, next(iter({v[1] for v in assigned})) if len({v[1] for v in assigned}) == 1 else None
            issue('负责人待部门确认', code, owners[code])
        if any(ord(c) < 32 for c in code) or code != code.strip():
            issue('产品编码格式', code, {'来源标题行': [b['start_row'] for b in blocks], '说明': '不自动删除换行或合并编码'})
        lines = []
        for line in block['lines']:
            if line['excluded']:
                continue
            ref = line['normalized_feed']
            quantity = Decimal(line['quantity'])
            recipe_ref = ref in groups and ref not in ERP_CONCENTRATIONS
            candidate_line = {'code': line['feed'], 'ref': ids[ref] if recipe_ref else ref,
                          'kind': 'recipe' if recipe_ref else 'material', 'quantity': line['quantity'],
                          'ratio': str(Decimal(line['ratio']) * 100), 'source_row': line['source_row']}
            input_key = (code, line['source_row'])
            if input_key in stock_inputs:
                choice = stock_inputs[input_key]
                assert choice['feed'] == ref and Decimal(choice['quantity']) == quantity
                assert choice['price_basis'] == 'erp_inventory' and choice['confirmed_on']
                candidate_line.update(kind='material', ref=ref, inventory_input=True, input_role='已有成品或返料掺料')
                applied_stock_inputs.add(input_key)
            if ref in replacements:
                choice = replacements[ref]
                assert candidate_line['kind'] == 'material' and choice['confirmed_on']
                assert choice['relationship'] in {'approved_substitute', 'same_material_different_manufacturer', 'confirmed_price_reference', 'confirmed_price_mean'}
                if choice['relationship'] == 'confirmed_price_mean':
                    assert choice['mean_policy'] == 'arithmetic_available_price'
                    targets = choice['reference_codes']
                    assert len(targets) >= 2 and len(set(targets)) == len(targets)
                    candidate_line['price_refs'] = targets
                    if quantity:
                        refs.update(targets)
                elif choice['relationship'] == 'confirmed_price_reference':
                    candidate_line['price_ref'] = choice['replacement']
                    if quantity:
                        refs.add(choice['replacement'])
                else:
                    assert choice['replacement'] not in groups
                    candidate_line['ref'] = choice['replacement']
                applied_replacements[ref] += 1
            if quantity:
                refs.add(candidate_line['ref'] if candidate_line['kind'] == 'material' else ref)
            if candidate_line['kind'] == 'material' and candidate_line['ref'] in zero_inputs:
                assert zero_inputs[candidate_line['ref']]['confirmed_on']
                candidate_line['unpriced_cost_zero'] = True
                applied_zero_inputs[candidate_line['ref']] += 1
            lines.append(candidate_line)
        mes = yields.get(code)
        yield_value = Decimal(mes['mean']) if mes else sum((Decimal(l['ratio']) / 100 for l in lines), Decimal(0))
        yield_source = {'kind': 'mes_mean' if mes else 'bom_ratio_sum', 'count': mes['orders'] if mes else 0}
        original_yield_value, original_yield_source = str(yield_value), deepcopy(yield_source)
        if code in yield_selections:
            choice = yield_selections[code]
            assert choice['source_sha256'] == hashes['yields']
            source_yields = list(db.execute('select source_row,yield_value from mes_yields where product=? order by source_row', (code,)))
            assert sorted(choice['source_rows']) == [r[0] for r in source_yields]
            mode = choice.get('mode', 'fixed_value')
            if mode in {'mes_mean', 'mes_mean_after_exclusions'}:
                excluded = set(choice['excluded_source_rows'])
                assert bool(excluded) == (mode == 'mes_mean_after_exclusions')
                assert len(excluded) == len(choice['excluded_source_rows']) and excluded < set(choice['source_rows'])
                included = [(n, Decimal(value)) for n,value in source_yields if n not in excluded]
                selected_yield = sum(value for _,value in included) / len(included)
                yield_source = dict(yield_source, count=len(included), source_rows=[n for n,_ in included], excluded_source_rows=sorted(excluded), original_count=len(source_yields), original_yield=str(yield_value), confirmed_on=choice['confirmed_on'])
            elif mode == 'bom_ratio_after_mes_exclusions':
                assert source_yields and sorted(choice['excluded_source_rows']) == [n for n,_ in source_yields]
                assert choice['bom_sha256'] == hashes['bom'] and choice['bom_header_row'] == block['start_row']
                selected_yield = sum((Decimal(l['ratio']) / 100 for l in lines), Decimal(0))
                yield_source = dict(yield_source, kind='bom_ratio_sum', count=0, source_rows=[], excluded_source_rows=sorted(choice['excluded_source_rows']), original_kind=yield_source['kind'], original_count=len(source_yields), original_yield=str(yield_value), confirmed_on=choice['confirmed_on'])
            else:
                assert mode == 'fixed_value'
                selected_yield = Decimal(choice['yield'])
                if 'confirmed_bom_ratio_percent_by_row' in choice:
                    assert not mes and not source_yields
                    assert choice['bom_sha256'] == hashes['bom'] and choice['bom_header_row'] == block['start_row']
                    assert choice['bom_detail_rows'] == [l['source_row'] for l in lines]
                    confirmed = choice['confirmed_bom_ratio_percent_by_row']
                    assert set(confirmed) == {str(l['source_row']) for l in lines}
                    interpretation = choice['bom_ratio_interpretation']
                    if interpretation == 'input_total':
                        total = sum((Decimal(l['quantity']) for l in lines), Decimal(0))
                        assert total.is_finite() and total > 0 and total == Decimal(choice['bom_ratio_denominator'])
                        ratios = [Decimal(l['quantity']) / total * 100 for l in lines[:-1]]
                        ratios.append(Decimal(100) - sum(ratios, Decimal(0)))
                    elif interpretation == 'percent_number':
                        assert Decimal(choice['bom_ratio_scale']) == Decimal('.01')
                        ratios = [Decimal(l['ratio']) * Decimal('.01') for l in lines]
                    else:
                        assert interpretation == 'original_ratio'
                        ratios = [Decimal(l['ratio']) for l in lines]
                    assert sum(ratios, Decimal(0)) / 100 == selected_yield
                    for line, ratio in zip(lines, ratios):
                        assert ratio.is_finite() and ratio == Decimal(confirmed[str(line['source_row'])])
                        assert Decimal(line['quantity']) != 0 or ratio == 0
                        line['ratio'] = decimal_text(ratio)
                yield_source = dict(yield_source, kind='user_confirmed', source_rows=choice['source_rows'], original_kind=yield_source['kind'], original_yield=str(yield_value), confirmed_on=choice['confirmed_on'])
            assert selected_yield.is_finite() and selected_yield > 0 and selected_yield * 100 == Decimal(choice['yield_percent'])
            assert selected_yield == Decimal(choice['yield'])
            yield_value = selected_yield
            applied_yield_selections.add(code)
        if not mes and code not in yield_selections and (yield_value < Decimal('.8') or yield_value > Decimal('1.2') or yield_value <= 0):
            issue('无MES的替代收率异常', code, {'标题行': block['start_row'], '比例总和': str(yield_value), '明细': [{'feed': l['code'], 'quantity': l['quantity'], 'ratio': str(Decimal(l['ratio']) / 100)} for l in lines]})
        formula = {'id': ids[code], 'name': code, 'kind': 'recipe', 'owner': assignment[0], 'department': assignment[1],
                   'source_kind': 'erp_bom', 'source_row': block['start_row'], 'source_hashes': hashes,
                   'ratio_linked': False, 'yield': str(yield_value),
                   'yield_source': yield_source,
                   'lines': lines}
        resolved[code] = formula
        if code in system_selections:
            # ERP配方采用与后续MES收率确认分别校验；保持原批准的整份来源指纹。
            erp_adoption_formulas[code] = deepcopy({**formula, 'yield': original_yield_value, 'yield_source': original_yield_source})
        db.execute('insert into resolved_formulas values(?,?,?,?)', (code, packed(formula), packed([b['start_row'] for b in blocks]), basis))

    assert applied_stock_inputs == stock_inputs.keys(), '已确认库存掺料未完整匹配原BOM行'
    assert applied_replacements.keys() == replacements.keys(), '已确认替代关系未匹配BOM'
    assert applied_zero_inputs.keys() == zero_inputs.keys(), '缺价按零核算口径未匹配投料'
    assert applied_yield_selections == yield_selections.keys(), '已确认收率未完整匹配产品'
    rules['confirmed_stock_inputs'] = list(stock_inputs.values())
    rules['confirmed_material_substitutions'] = replacements
    rules['unpriced_zero_cost_inputs'] = zero_inputs
    rules['confirmed_yield_selections'] = yield_selections
    db.execute('update confirmed_rules set payload_json=?', (packed(rules),))
    refs.add('CF232')
    samples = defaultdict(list)
    for r in db.execute('select * from mes_yields order by source_row'):
        samples[r['product']].append((r['source_row'], r['yield_value']))
    for code, mes in yields.items():
        if code not in groups:
            issue('MES无对应BOM', code, {'工单行': samples[code], '平均收率': mes['mean'], '差异': mes['spread']})
            continue
        reasons = []
        if Decimal(mes['spread']) > Decimal('.05'):
            reasons.append('最大最小差超过5个百分点')
        if Decimal(mes['minimum']) < Decimal('.8') or Decimal(mes['maximum']) > Decimal('1.2'):
            reasons.append('包含低于80%或高于120%的筛查值')
        if reasons and code not in yield_selections:
            issue('MES收率待确认', code, {'原因': reasons, '部门': resolved[code]['department'], '工单行': samples[code], '平均收率': mes['mean'], '差异': mes['spread'], '说明': '不自动删除、去重或截断工单'})
    for code in owners.keys() - groups.keys():
        issue('负责人无对应BOM', code, owners[code])

    price_items = {str(r[8]): (n, r) for n, r in current['erp_prices']}
    price_names = defaultdict(list)
    for n, r in current['erp_prices']:
        assert r[12] is not None and Decimal(str(r[12])) > 0
        price_names[r[9]].append((n, r))
    scope_rows = [(n, r) for n, r in current['erp_stock'] if r[5] in refs and r[8] in WAREHOUSES]
    stock_names = defaultdict(list)
    cross_rows, no_price_items = [], defaultdict(list)
    for n, row in scope_rows:
        stock_names[row[5]].append((n, row))
        match = price_items.get(str(row[4]))
        pn, price = match if match else (None, None)
        cross = bool(price and (str(row[0]), str(row[2])) != (str(price[2]), str(price[4])))
        if cross:
            cross_rows.append({'库存行': n, '品号': row[4], '品名': row[5], '库存公司': row[1], '库存工厂': row[3], '单价公司': price[3], '单价工厂': price[5], '单价': price[12]})
        if not price:
            no_price_items[str(row[4])].append({'库存行': n, '品名': row[5], '公司': row[1], '仓库': row[8], '数量': row[9]})
        db.execute('insert into inventory_rows values(?,?,?,?,?,?,?,?,?,?,?,?,?)',
                   (n, row[5], str(row[4]), row[1], row[3], row[8], str(row[9]), row[7],
                    str(price[12]) if price else None, pn, price[3] if price else None, price[5] if price else None, int(cross)))
    if cross_rows and not shared_prices:
        issue('跨公司或工厂取价口径', '共同口径', {'涉及库存行数': len(cross_rows), '示例': cross_rows[:8], '说明': '试算按唯一品号匹配，不代表已批准跨公司共用单价'})
    if not cutoff:
        issue('库存数据实际截止日期', '共同口径', {'数量截止日期': None, '单价记录日期': '逐行保留创建日期和会计期间，不能替代数量截止日期'})
    baseline_prices = live['prices']
    prices = {}
    inventory_export = []
    for code in sorted(refs):
        rows = stock_names.get(code, [])
        units = {r[7] for _, r in rows}
        assert len(units) <= 1, f'库存混合单位：{code}'
        quantity = sum((Decimal(str(r[9])) for _, r in rows), Decimal(0)) if rows else None
        if quantity is not None and quantity < 0:
            issue('范围内负库存', code, {'来源行': [n for n, r in rows], '数量': str(quantity)})
        item_ids = {str(r[4]) for _, r in rows}
        matches = [price_items[k] for k in sorted(item_ids) if k in price_items] if item_ids else price_names.get(code, [])
        values = {Decimal(str(r[12])) for _, r in matches}
        stock_price = next(iter(values)) if len(values) == 1 and (not item_ids or item_ids <= price_items.keys()) else None
        if len(values) > 1:
            issue('同名多品号取价', code, {'库存品号': sorted(item_ids), '价格记录': [{'行号': n, '品号': r[8], '单价': r[12]} for n, r in matches]})
        latest = baseline_prices.get(code, {}).get('latest_price')
        prices[code] = {'latest_price': latest, 'inventory_price': decimal_text(stock_price)}
        evidence = {'库存行': [n for n, r in rows], '价格记录': [{'行号': n, '品号': r[8], '公司': r[3], '工厂': r[5], '年度': r[6], '期间': r[7], '创建日期': r[13], '单位成本': r[12]} for n, r in matches], '负责人处理': '沿用现有负责人' if code in baseline_prices else 'ERP新增负责人留空', '来源': 'ERP 库存', '库存数量截止日期': cutoff}
        unit = next(iter(units)) if units else None
        db.execute('insert into material_candidates values(?,?,?,?,?,?)', (code, unit, decimal_text(quantity), decimal_text(stock_price), latest, packed(evidence)))
        inventory_export.append([code, '／'.join(sorted(item_ids)), unit or '', decimal_text(quantity) if quantity is not None else '', decimal_text(stock_price) if stock_price is not None else '', latest or '', 'ERP 库存', cutoff or '', packed(evidence)])
    prices['纯水'] = {'latest_price': '0', 'inventory_price': None}
    for code, entry in baseline_prices.items():
        prices.setdefault(code, {'latest_price': entry.get('latest_price'), 'inventory_price': None})

    package = deepcopy(live['package'])
    package['recipes'] = [r for r in live['formulas'] if r['name'] not in resolved] + list(resolved.values())
    fallback_codes = sorted(code for code in resolved if Decimal(prices.get(code, {}).get('inventory_price') or '0') > 0)
    package['erp_inventory_fallback_codes'] = fallback_codes
    package.setdefault('current_policy', {})['erp_inventory_fallback_codes'] = fallback_codes
    recipe_by_id = {r['id']: r for r in package['recipes']}
    assert len(recipe_by_id) == len(package['recipes'])
    graph = {r['id']: {l['ref'] for l in r['lines'] if l['kind'] != 'material' and Decimal(l['quantity']) != 0} for r in package['recipes']}
    state, stack, cycles = {}, [], []

    def visit(key):
        state[key] = 1
        stack.append(key)
        for ref in sorted(graph[key]):
            assert ref in graph, f'引用无配方：{ref}'
            if state.get(ref) == 1:
                cycles.append(stack[stack.index(ref):] + [ref])
            elif not state.get(ref):
                visit(ref)
        stack.pop()
        state[key] = 2

    for key in sorted(graph):
        if not state.get(key):
            visit(key)
    blocked = {k for cycle in cycles for k in cycle}
    changed = True
    while changed:
        earlier = len(blocked)
        blocked.update(k for k, refs_ in graph.items() if refs_ & blocked)
        changed = len(blocked) != earlier
    for cycle in cycles:
        issue('循环投料口径', recipe_by_id[cycle[0]]['name'], [recipe_by_id[k]['name'] for k in cycle])
    safe = deepcopy(package)
    safe['recipes'] = [r for r in package['recipes'] if r['id'] not in blocked]
    latest_results = calculate(safe, prices, price_policy='latest')
    erp_only = {code: dict(entry, latest_price='0' if code == '纯水' else None) for code, entry in prices.items()}
    inventory_results = calculate(safe, erp_only, price_policy='inventory')
    cost_csv = []
    missing_uses = defaultdict(lambda: {'latest': set(), 'inventory': set()})
    for code, formula in sorted(resolved.items()):
        records = []
        for policy, results in [('latest', latest_results), ('inventory', inventory_results)]:
            if formula['id'] in blocked:
                result, status = None, '循环引用待确认'
            else:
                result = results[formula['id']]
                status = '缺价' if result['cost'] is None else '试算可计算；业务口径待确认'
                for ref in result['missing_materials']:
                    missing_uses[ref][policy].add(code)
                if result['cost'] is not None:
                    amount = sum((Decimal(l['amount']) for l in result['lines'] if l['amount'] is not None), Decimal(0))
                    expected = amount / Decimal(result['total_input']) / Decimal(result['yield'])
                    assert abs(expected - Decimal(result['cost'])) < Decimal('1e-20')
            db.execute('insert into trial_costs values(?,?,?,?,?,?)', (code, policy, result['cost'] if result else None, packed(result['missing_materials'] if result else []), status, packed(result)))
            records.append(result)
        cost_csv.append([code, formula['department'] or '', formula['owner'] or '', formula['yield'], formula['yield_source']['kind'],
                         records[0]['cost'] if records[0] else '', records[1]['cost'] if records[1] else '',
                         '循环引用待确认' if formula['id'] in blocked else packed({'最新缺价': records[0]['missing_materials'], 'ERP库存缺价': records[1]['missing_materials']})])
    direct = {ref for f in resolved.values() for l in f['lines'] if l['kind'] == 'material' and Decimal(l['quantity']) != 0 for ref in price_codes(l)} - {'纯水'}
    unresolved_direct = sorted(r for r in direct if r not in zero_inputs and prices.get(r, {}).get('latest_price') is None and prices.get(r, {}).get('inventory_price') is None)
    for ref in unresolved_direct:
        uses = sorted(f['name'] for f in resolved.values() if any(l['kind'] == 'material' and ref in price_codes(l) and Decimal(l['quantity']) != 0 for l in f['lines']))
        issue('直接投料仍无可用价格', ref, {'直接引用产品': uses, '涉及部门': sorted({resolved[p]['department'] for p in uses if resolved[p]['department']}), '相似编码待人工核对': difflib.get_close_matches(ref, sorted(price_names.keys() | baseline_prices.keys()), n=3, cutoff=.8), '说明': '未用旧库存补价或零价填充；相似名称不自动建立对应关系'})
    for item, rows in sorted(no_price_items.items()):
        code = rows[0]['品名']
        result = inventory_results.get(ids.get(code))
        issue('库存品号缺ERP单价', code, {'品号': item, '库存明细': rows, 'BOM试算成本': result['cost'] if result else None, '说明': 'BOM可算成本不能自动当作ERP原始库存单价'})

    difference_counts = Counter()
    difference_csv = []
    approved_system_csv = []
    for old_formula in live['formulas']:
        if old_formula['kind'] != 'recipe':
            continue
        code = old_formula['name']
        if code not in resolved:
            evidence = {'说明': '按已确认规则保留现役产品，不删除', '原配方': old_formula}
            difference_counts['保留现有产品'] += 1
            db.execute('insert into system_differences values(?,?,?)', (code, '保留现有产品', packed(evidence)))
            difference_csv.append([code, '保留现有产品', old_formula['yield'], '', '已确认保留，不列为待决问题', '已确认保留现役产品'])
            continue
        new = resolved[code]
        left, right = summarize_inputs(old_formula), summarize_inputs(new)
        lt, rt = sum(left.values(), Decimal(0)), sum(right.values(), Decimal(0))
        if left == right:
            classification = '汇总投料一致'
        elif lt > 0 and rt > 0 and all(left.get(k, Decimal(0)) * rt == right.get(k, Decimal(0)) * lt for k in left.keys() | right.keys()):
            classification = '仅整批量等比变化'
        else:
            classification = '化工投料实质变化'
        difference_counts[classification] += 1
        evidence = {'原配方': old_formula, '候选配方': new, '原料汇总差异': [{'code': k, 'old_quantity': str(left.get(k, Decimal(0))), 'new_quantity': str(right.get(k, Decimal(0)))} for k in sorted(left.keys() | right.keys()) if left.get(k, Decimal(0)) != right.get(k, Decimal(0))], '原收率': old_formula['yield'], '新收率': new['yield']}
        decision = ''
        if code in system_selections:
            assert classification == '化工投料实质变化' and system_confirmation['choice'] == 'erp'
            assert new['department'] == system_confirmation['department']
            assert system_selections[code] == {'old_formula_sha256': formula_digest(old_formula), 'erp_formula_sha256': formula_digest(erp_adoption_formulas[code])}, f'已确认配方发生变化，重新核对：{code}'
            applied_system_selections.add(code)
            decision = '已确认采用ERP配方；收率和缺价疑点另行核对'
            evidence['配方采用确认'] = dict(system_selections[code], choice='erp', confirmed_on=system_confirmation['confirmed_on'])
            approved_system_csv.append([code, new['department'], new['owner'], new['source_row'], 'ERP配方', system_confirmation['confirmed_on']])
        db.execute('insert into system_differences values(?,?,?)', (code, classification, packed(evidence)))
        difference_csv.append([code, classification, old_formula['yield'], new['yield'], packed(evidence['原料汇总差异']), decision])
        if classification == '化工投料实质变化' and not decision:
            issue('现役产品配方替换范围', code, evidence)

    assert applied_system_selections == system_selections.keys(), '已确认现役配方未完整匹配'
    rules['confirmed_system_formula_replacements'] = system_confirmation
    db.execute('update confirmed_rules set payload_json=?', (packed(rules),))
    checks = {
        'source_hashes': len(hashes), 'raw_rows': db.execute('select count(*) from raw_rows').fetchone()[0],
        'resolved_products': len(resolved), 'cycle_groups': len(cycles), 'cycle_affected_products': sum(f['id'] in blocked for f in resolved.values()),
        'confirmed_formula_choices': len(decisions['formula_selections']), 'confirmed_owner_choices': len(decisions['owner_selections']),
        'confirmed_stock_input_rows': len(applied_stock_inputs),
        'confirmed_material_substitution_rows': dict(applied_replacements),
        'confirmed_unpriced_zero_input_rows': dict(applied_zero_inputs),
        'confirmed_system_formula_replacements': len(applied_system_selections),
        'confirmed_yield_selections': len(applied_yield_selections),
        'same_item_multiplied_price_rows': 0, 'stock_grain_duplicates': 0,
    }
    for r in db.execute('select * from sources'):
        assert digest(originals / r['filename']) == r['sha256']
        assert db.execute('select count(*) from raw_rows where kind=?', (r['kind'],)).fetchone()[0] == r['row_count']
    assert resolved['T21-DB']['source_row'] == 13173
    assert resolved['B-71D']['source_row'] == 12709 and resolved['B-71B']['source_row'] == 17642
    assert resolved['B-9D']['owner'] == '吴振宇' and resolved['A-85KJ']['owner'] == '李树荣'
    if 'A-111' in decisions['owner_selections']:
        assert resolved['A-111']['owner'] == decisions['owner_selections']['A-111']['owner']
    else:
        assert resolved['A-111']['owner'] is None
    assert prices['T2-D']['inventory_price'] == '3.535706'
    assert all(row[1] not in {'宏昊虚拟仓', '宏图虚拟仓'} for row in db.execute('select code,warehouse from inventory_rows'))
    assert abs(sum(Decimal(l['ratio']) for l in resolved['GW96-01A']['lines']) - 100) < Decimal('1e-20')
    negative = [l for f in resolved.values() for l in f['lines'] if Decimal(l['quantity']) < 0]
    assert db.execute('select count(*) from bom_lines where cast(quantity as real)<0').fetchone()[0] == 34
    assert negative and all(l['kind'] == 'material' and l['ref'] in {'回收CF232', 'CF232D'} for l in negative)
    db.commit()
    assert db.execute('pragma integrity_check').fetchone()[0] == 'ok'
    cost_counts = {policy: Counter('循环引用待确认' if r['status'] == '循环引用待确认' else '可计算' if r['cost'] is not None else '缺价' for r in db.execute('select * from trial_costs where policy=?', (policy,))) for policy in ['latest', 'inventory']}
    summary = {'created_at': datetime.now(timezone(timedelta(hours=8))).isoformat(), 'database_write_scope': 'independent_analysis_only',
               'source_rows': {r['kind']: r['row_count'] for r in db.execute('select * from sources')}, 'checks': checks,
               'inventory_scope_rows': len(scope_rows), 'inventory_scope_codes': len(stock_names), 'cross_entity_rows': len(cross_rows),
               'no_erp_price_stock_items': len(no_price_items), 'unpriced_direct_inputs': len(unresolved_direct),
               'identified_bom_inventory_fallback_codes': len(fallback_codes),
               'trial_costs': cost_counts, 'system_difference_classes': difference_counts,
               'issues': Counter(c for c, p, e in issues),
               'inventory_confirmation': inventory_confirmation,
               'assumptions': ['已确认同一ERP品号跨公司和工厂统一取价' if shared_prices else '试算暂按唯一品号对应单价，跨公司或工厂共享价格尚未批准', '最新成本保留现役最新采购价，缺最新价时仅用新ERP价格', '库存成本试算仅用新ERP价格，纯水按已确认零成本；未采用旧库存或最新采购价补空', '原始MES和BOM比例全部保留；候选按逐项确认的收率及比例处理，未确认项继续待核对', '等价配方代表段仅用于分析，原始配方段和行号全部保留']}
    summary['unpriced_zero_cost_inputs'] = zero_inputs
    if zero_inputs:
        summary['assumptions'].append('用户确认的催化剂缺价按0核算，原始价格保持空；保留产品和投料量，不标为实际采购零价')
    (output / '分析摘要.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    (output / '已确认规则.json').write_text(json.dumps(rules, ensure_ascii=False, indent=2), encoding='utf-8')
    (output / '双身份库存兜底编码.json').write_text(json.dumps(fallback_codes, ensure_ascii=False, indent=2), encoding='utf-8')
    (output / '成本试算输入.json').write_text(json.dumps({'package': package, 'prices': prices, 'cycle_blocked_recipe_ids': sorted(blocked)}, ensure_ascii=False, indent=2), encoding='utf-8')
    shutil.copy2(decisions_path, output / decisions_path.name)
    shutil.copy2(live_path, output / '现有系统只读基线-20261007.json')
    write_csv(output / '库存匹配清单.csv', ['原料编码', '库存品号', '单位', '范围内库存数量', '新ERP库存单价', '现役最新采购价', '来源', '库存数量实际截止日期', '来源证据'], inventory_export)
    write_csv(output / '产品双成本试算.csv', ['产品', '部门', '负责人', '收率原值', '收率来源', '最新成本试算', 'ERP库存成本试算', '待核对信息'], cost_csv)
    write_csv(output / '现役配方差异.csv', ['产品', '差异类型', '现役收率', '候选收率', '投料差异', '配方采用确认'], difference_csv)
    write_csv(output / '已确认现役配方采用清单.csv', ['产品', '部门', '负责人', 'ERP BOM标题行', '采用配方', '确认日期'], approved_system_csv)
    write_csv(output / '已确认收率清单.csv', ['产品', '部门', '负责人', '确认收率%', '原始收率%', 'MES原始行号', '确认日期', '确认说明'], [[code, resolved[code]['department'], resolved[code]['owner'], choice['yield_percent'], str(Decimal(resolved[code]['yield_source']['original_yield']) * 100), packed(choice['source_rows']), choice['confirmed_on'], choice['scope']] for code, choice in yield_selections.items()])
    write_csv(output / '跨公司取价明细.csv', ['库存行', '品号', '品名', '库存公司', '库存工厂', '单价公司', '单价工厂', '单价'], [[r[k] for k in ['库存行', '品号', '品名', '库存公司', '库存工厂', '单价公司', '单价工厂', '单价']] for r in cross_rows])
    write_csv(output / '待确认清单.csv', ['序号', '类别', '编码', '证据', '确认结果'], [[i, c, p, packed(e), ''] for i, (c, p, e) in enumerate(issues, 1)])
    write_csv(output / '待补价及编码对应.csv', ['投料编码', '直接引用产品数', '涉及部门', '直接引用产品', '相似编码仅供核对', '确认对应编码', '确认ERP库存单价', '确认最新采购单价', '确认说明'], [[p, len(e['直接引用产品']), '／'.join(e['涉及部门']), '／'.join(e['直接引用产品']), '／'.join(e['相似编码待人工核对']), '', '', '', ''] for c, p, e in issues if c == '直接投料仍无可用价格'])
    mes_review = [(p, e) for c, p, e in issues if c == 'MES收率待确认']
    write_csv(output / 'MES收率待确认.csv', ['产品', '部门', '工单数', '平均收率%', '最低收率%', '最高收率%', '最大最小差百分点', '核对原因', '原表行号及收率', '业务确认'], [[p, e['部门'], yields[p]['orders'], str(Decimal(yields[p]['mean']) * 100), str(Decimal(yields[p]['minimum']) * 100), str(Decimal(yields[p]['maximum']) * 100), str(Decimal(yields[p]['spread']) * 100), '；'.join(e['原因']), packed(e['工单行']), ''] for p, e in mes_review])
    previous_page = (prior / '研发数据核对.html').read_text(encoding='utf-8')
    sections = []
    for category in summary['issues']:
        entries = [(i, p, e) for i, (c, p, e) in enumerate(issues, 1) if c == category]
        body = ''.join(f'<details><summary>{i}. {html.escape(p)}</summary>{render_evidence(e)}</details>' for i, p, e in entries)
        sections.append(f'<section><h2>{html.escape(category)}（{len(entries)}）</h2>{body}</section>')
    before, after = previous_page.split('<main>', 1)
    _, tail = after.split('</main>', 1)
    owner_pending = summary['issues'].get('负责人待部门确认', 0)
    intro = f'<section><h2>{summary["created_at"][:10]} 核对进度</h2><p>2380 个产品已形成分析候选；20 组配方选择已应用，负责人待确认 {owner_pending} 项，现役差异配方已确认采用ERP {len(applied_system_selections)} 项。库存明细和单价补充表已按品号匹配。以下成本均为待业务确认的试算，未写入正式库。</p><p><a href="产品双成本试算.csv">双成本试算</a> · <a href="MES收率待确认.csv">MES收率核对</a> · <a href="库存匹配清单.csv">库存匹配</a> · <a href="跨公司取价明细.csv">跨公司取价明细</a> · <a href="现役配方差异.csv">现役配方差异</a> · <a href="已确认现役配方采用清单.csv">已确认现役配方采用清单</a></p><pre>{html.escape(json.dumps(summary['trial_costs'], ensure_ascii=False, indent=2))}</pre></section>'
    (output / '研发数据核对.html').write_text(before + intro + '<main>' + ''.join(sections) + '</main>' + tail, encoding='utf-8')
    (output / '阅读说明.txt').write_text(f'本目录为独立分析候选，不是正式导入结果。\n先打开研发数据核对.html查看合并后的业务问题。\n来源原件保留{len(hashes)}份来源；分析DB保留原值、公式、行号、哈希及旧批次问题。\n业务确认记录20项配方选择、{len(decisions["owner_selections"])}项负责人选择，待部门确认{owner_pending}项。\n产品双成本试算仅使用现役最新采购价与新ERP价格；未用旧库存补价或零填充。\n库存数量截止日期：{cutoff or "待确认"}；跨公司或工厂按同品号统一取价：{"已确认" if shared_prices else "待确认"}。\n已确认库存掺料{len(applied_stock_inputs)}行，仅指定行在双成本中固定使用ERP库存价；原BOM编码、投料量和顺序保留。\n双身份编码按已确认规则先算BOM，缺价时用对应ERP库存价；不掩盖未经确认的循环引用。\n无MES时默认使用化工明细比例合计；已确认项按逐项决定更新候选比例与收率，原始BOM保持原值。\n剩余问题见待确认清单，异常收率及其他未确认事项继续核对。\n全部问题确认后才准备统一正式写入，当前未提交、推送或部署。\n', encoding='utf-8')
    db.close()
    return summary


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', required=True, type=Path)
    parser.add_argument('--prior', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--live-snapshot', required=True, type=Path)
    args = parser.parse_args()
    print(packed(reconcile(args.source, args.prior, args.output, args.live_snapshot)))
