"""Source-preserving research staging; never writes to the application database."""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import re
import sqlite3
from collections import Counter, defaultdict
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


HEADERS = {
    "bom": ["产品内编", "投料编码", "投料量", "配方比例", "计算底数", "成本计算标记"],
    "owners": ["产品内编", "负责人", "部门"],
    "yields": ["品名", "实际产量", "总投料量", "预计用量", "水总量", "掺料", "产品收率"],
    "inventory_reference": ["品名", "库存数量", "成本金额", "单位成本"],
}
RULES = {
    "inventory_status": "等待新库存表；本次库存只作参考，禁止正式导入",
    "mes_rows": "每行作为一个工单，不按相同值去重",
    "yield_spread": "0.05",
    "yield_without_mes": "不含标题及包装物的明细比例之和",
    "packaging": "保留原始记录，排除成本、净投料与收率",
    "water": "水及W系列水引用纯水",
    "recovery": "负投料抵减成本与净投料；回收2引用回收CF232",
    "concentrations": {"CF232A": "0.25", "CF232B": "0.65", "CF232D": "0.25", "回收CF232": "0.90"},
    "recipe_fallback": "已识别双重身份编码先算BOM，缺价时用ERP库存；循环须先确认",
}


def packed(value):
    return json.dumps(value, ensure_ascii=False, default=str)


def packaging(code):
    return any(word in code for word in ("包材", "胶袋", "纸箱", "胶桶", "铁桶", "标签"))


def normalized_feed(code):
    if code == "水" or re.fullmatch(r"W\d+水", code):
        return "纯水"
    return "回收CF232" if code == "回收2" else code


def render_evidence(evidence):
    def text(value):
        return html.escape(str(value))
    def table(headers, rows):
        return '<table><thead><tr>' + ''.join('<th>'+text(h)+'</th>' for h in headers) + '</tr></thead><tbody>' + ''.join('<tr>'+''.join('<td>'+text(v)+'</td>' for v in row)+'</tr>' for row in rows) + '</tbody></table>'
    if isinstance(evidence, list):
        if evidence and isinstance(evidence[0], str):
            return '<p>'+text(' → '.join(evidence))+'</p>'
        if evidence and isinstance(evidence[0], dict) and 'owner' in evidence[0]:
            return table(['负责人','部门','原表行号'], [(r['owner'],r['department'],r.get('row',r.get('source_row',''))) for r in evidence])
        if evidence and isinstance(evidence[0], dict) and 'lines' in evidence[0]:
            return ''.join('<h3>配方段 · 标题行 '+text(b['row'])+'</h3>'+table(['投料编码','投料量','明细比例','参与核算','原表行号'], [(l['feed'],l['quantity'],str(Decimal(l['ratio'])*100)+'%','否（包装物）' if l['excluded'] else '是',l['row']) for l in b['lines']]) for b in evidence)
    if isinstance(evidence, dict):
        if '工单行' in evidence:
            return '<p>平均收率：'+text(str(Decimal(evidence['平均收率'])*100))+'%；最大最小差：'+text(str(Decimal(evidence['差异'])*100))+' 个百分点</p>'+table(['原表行号','收率'], [(n,str(Decimal(v)*100)+'%') for n,v in evidence['工单行']])
        if '原配方' in evidence:
            old=evidence['原配方']
            content='<h3>系统现有配方</h3><p>负责人：'+text(old.get('owner',''))+'</p>'+table(['投料编码','投料量'],[(l['code'],l['quantity']) for l in old['lines']])
            new_blocks=evidence.get('新BOM',[])
            if '候选配方' in evidence:
                new=evidence['候选配方']
                new_blocks=[{'row':new['source_row'],'lines':[{'feed':l['code'],'quantity':l['quantity'],'ratio':str(Decimal(l['ratio'])/100),'excluded':False,'row':l['source_row']} for l in new['lines']]}]
            return content+render_evidence(new_blocks)+render_evidence(evidence.get('新归属',[]))
        if '比例总和' in evidence:
            return '<p>拟用收率：'+text(str(Decimal(evidence['比例总和'])*100))+'%；标题行：'+text(evidence['标题行'])+'</p>'+table(['投料编码','投料量','明细比例'],[(l['feed'],l['quantity'],str(Decimal(l['ratio'])*100)+'%') for l in evidence['明细']])
        if '引用产品' in evidence:
            content='<p>引用产品：'+text('、'.join(evidence['引用产品']))+'</p>'
            if evidence.get('旧库存参考'):
                content+='<p>旧库存仅供比较</p>'+table(HEADERS['inventory_reference'],[evidence['旧库存参考']])
            return content
    return '<pre>'+html.escape(json.dumps(evidence,ensure_ascii=False,indent=2,default=str))+'</pre>'


def analyze(source: Path, output: Path, snapshot: dict | None = None):
    output.mkdir(parents=True, exist_ok=True)
    target = output / "研发数据分析.db"
    if target.exists():
        raise ValueError("分析库已存在；请使用新的输出目录以保留上一批证据")
    db = sqlite3.connect(target)
    db.executescript("""
      CREATE TABLE sources(kind TEXT PRIMARY KEY,filename TEXT,sha256 TEXT,sheet TEXT,header_json TEXT,row_count INTEGER);
      CREATE TABLE raw_rows(kind TEXT,source_row INTEGER,values_json TEXT,formula_json TEXT,PRIMARY KEY(kind,source_row));
      CREATE TABLE bom_blocks(block_id INTEGER PRIMARY KEY,product TEXT,start_row INTEGER);
      CREATE TABLE bom_lines(block_id INTEGER,source_row INTEGER,feed TEXT,normalized_feed TEXT,quantity TEXT,ratio TEXT,excluded INTEGER);
      CREATE TABLE product_owners(product TEXT,owner TEXT,department TEXT,source_row INTEGER);
      CREATE TABLE mes_yields(product TEXT,yield_value TEXT,source_row INTEGER);
      CREATE TABLE yield_summary(product TEXT PRIMARY KEY,orders INTEGER,mean TEXT,minimum TEXT,maximum TEXT,spread TEXT);
      CREATE TABLE review_issues(issue_id INTEGER PRIMARY KEY,category TEXT,product TEXT,evidence_json TEXT,status TEXT DEFAULT '待确认',decision TEXT DEFAULT '');
      CREATE TABLE current_snapshot(kind TEXT,payload_json TEXT);
      CREATE TABLE confirmed_rules(payload_json TEXT);
      CREATE TABLE candidate_formulas(block_id INTEGER PRIMARY KEY,product TEXT,formula_json TEXT);
    """)
    db.execute("INSERT INTO confirmed_rules VALUES(?)", (packed(RULES),))
    rows = {}
    for file in sorted(source.rglob("*.xlsx")):
        if file.name.startswith("~$") or output.resolve() in file.resolve().parents or any(part.startswith("分析结果-") for part in file.relative_to(source).parts[:-1]):
            continue
        with file.open("rb") as stream:
            digest = hashlib.file_digest(stream, "sha256").hexdigest()
        values = load_workbook(file, read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(file, read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in values:
                data = list(sheet.iter_rows(values_only=True))
                header = list(data[0])
                kind = next((key for key, expected in HEADERS.items() if header == expected), None)
                if kind is None:
                    raise ValueError(f"未识别表头：{file.name}/{sheet.title}")
                if kind in rows:
                    raise ValueError(f"同批存在多份{kind}，请明确来源版本")
                raw = list(formulas[sheet.title].iter_rows(values_only=True))
                rows[kind] = [(n, list(row)) for n, row in enumerate(data[1:], 2) if any(v is not None for v in row)]
                db.execute("INSERT INTO sources VALUES(?,?,?,?,?,?)", (kind, file.name, digest, sheet.title, packed(header), len(rows[kind])))
                db.executemany("INSERT INTO raw_rows VALUES(?,?,?,?)", [(kind, n, packed(row), packed(raw[n-1])) for n, row in rows[kind]])
        finally:
            values.close()
            formulas.close()
    if not {"bom", "owners", "yields"} <= rows.keys():
        raise ValueError("缺少研发三表")
    issues = []
    def issue(category, product, evidence):
        issues.append((category, product, evidence))
        db.execute("INSERT INTO review_issues(category,product,evidence_json) VALUES(?,?,?)", (category, product, packed(evidence)))
    groups = defaultdict(list)
    block = None
    for n, row in rows["bom"]:
        product, feed, quantity, ratio, base, flag = row
        if flag == "否":
            block = {"id": n, "product": product, "row": n, "lines": []}
            groups[product].append(block)
            db.execute("INSERT INTO bom_blocks VALUES(?,?,?)", (n, product, n))
            if product != feed or base != 0 or quantity != 0:
                issue("标题异常", product, {"row": n, "values": row})
        elif flag != "是" or block is None or product != block["product"] or base != 1:
            issue("BOM分段异常", product, {"row": n, "values": row})
        else:
            line = {"row": n, "feed": feed, "ref": normalized_feed(feed), "quantity": str(quantity), "ratio": str(ratio), "excluded": packaging(feed)}
            block["lines"].append(line)
            db.execute("INSERT INTO bom_lines VALUES(?,?,?,?,?,?,?)", (block["id"], n, feed, line["ref"], str(quantity), str(ratio), int(line["excluded"])))
    owners = defaultdict(list)
    for n, row in rows["owners"]:
        owners[row[0]].append({"row": n, "owner": row[1], "department": row[2]})
        db.execute("INSERT INTO product_owners VALUES(?,?,?,?)", (*row, n))
    for code in sorted(groups):
        blocks = groups[code]
        if any(ord(c) < 32 for c in code) or code != code.strip():
            issue("产品编码格式待确认", code, {"标题行": [b['row'] for b in blocks], "说明": "不自行删除换行或空格"})
        signatures = {packed([{k: v for k, v in line.items() if k != "row"} for line in b["lines"]]) for b in blocks}
        if len(signatures) > 1:
            issue("不同BOM", code, blocks)
        if len({(r["owner"], r["department"]) for r in owners[code]}) > 1:
            issue("负责人冲突", code, owners[code])
        if not owners[code]:
            issue("缺少负责人", code, blocks)
    for code in sorted(owners.keys() - groups.keys()):
        issue("负责人无BOM", code, owners[code])
    yields = defaultdict(list)
    for n, row in rows["yields"]:
        value = Decimal(str(row[6]))
        yields[row[0]].append((n, value))
        db.execute("INSERT INTO mes_yields VALUES(?,?,?)", (row[0], str(value), n))
        if abs(value - Decimal(str(row[1])) / Decimal(str(row[2]))) > Decimal("0.000000001"):
            issue("MES计算不一致", row[0], {"row": n, "values": row})
    for code, samples in sorted(yields.items()):
        vals = [v for n, v in samples]
        mean, minimum, maximum = sum(vals) / len(vals), min(vals), max(vals)
        db.execute("INSERT INTO yield_summary VALUES(?,?,?,?,?,?)", (code, len(vals), str(mean), str(minimum), str(maximum), str(maximum-minimum)))
        evidence = {"工单行": [(n, str(v)) for n, v in samples], "平均收率": str(mean), "差异": str(maximum-minimum)}
        if maximum-minimum > Decimal("0.05"):
            issue("MES收率波动", code, evidence)
        if minimum <= 0 or minimum < Decimal("0.8") or maximum > Decimal("1.2"):
            issue("MES极端值待核对", code, evidence | {"说明": "80%/120%仅为筛查线，不自动剔除或截断"})
        if code not in groups:
            issue("MES无BOM", code, evidence)
    for code, blocks in sorted(groups.items()):
        for b in blocks:
            active = [l for l in b["lines"] if not l["excluded"]]
            ratios = sum((Decimal(l["ratio"]) for l in active), Decimal(0))
            if code not in yields and (ratios <= 0 or ratios < Decimal("0.8") or ratios > Decimal("1.2")):
                issue("替代收率异常", code, {"标题行": b["row"], "比例总和": str(ratios), "明细": active})
            for line in active:
                if Decimal(line["quantity"]) < 0 and line["ref"] not in {"回收CF232", "CF232D"}:
                    issue("未确认负投料", code, line)
    graph = {c: {l["ref"] for b in bs for l in b["lines"] if not l["excluded"] and Decimal(l["quantity"]) != 0 and l["ref"] in groups} for c, bs in groups.items()}
    state, path = {}, []
    def visit(code):
        state[code] = 1
        path.append(code)
        for ref in sorted(graph[code]):
            if state.get(ref) == 1:
                issue("循环引用", code, path[path.index(ref):] + [ref])
            elif not state.get(ref):
                visit(ref)
        path.pop()
        state[code] = 2
    for code in sorted(graph):
        if not state.get(code):
            visit(code)
    snapshot = snapshot or {}
    for kind, value in snapshot.items():
        db.execute("INSERT INTO current_snapshot VALUES(?,?)", (kind, packed(value)))
    current_prices = {r["code"]: r for r in snapshot.get("materials", [])}
    old_stock = {row[0]: row for n, row in rows.get("inventory_reference", [])}
    feeds = defaultdict(set)
    for code, bs in groups.items():
        for b in bs:
            for l in b["lines"]:
                ref = "CF232" if l["ref"] in RULES["concentrations"] else l["ref"]
                if not l["excluded"] and Decimal(l["quantity"]) != 0 and ref not in groups and ref != "纯水":
                    feeds[ref].add(code)
    for ref, users in sorted(feeds.items()):
        price = current_prices.get(ref, {})
        if price.get("latest_price") is None and price.get("inventory_price") is None:
            category = "待新库存补价" if ref in old_stock else "缺价或编码待对应"
            issue(category, ref, {"引用产品": sorted(users), "旧库存参考": old_stock.get(ref), "说明": "旧库存不作为本轮最终价格"})
    for formula in snapshot.get("formulas", []):
        code = formula["name"]
        if formula.get("kind") != "recipe":
            continue
        if code not in groups:
            issue("现有产品未在本批BOM", code, {"说明": "保留现有产品，不删除", "原配方": formula})
            continue
        old = sorted((normalized_feed(l["code"]), Decimal(l["quantity"])) for l in formula["lines"] if Decimal(l["quantity"]) != 0)
        new = [sorted((l["ref"], Decimal(l["quantity"])) for l in b["lines"] if not l["excluded"] and Decimal(l["quantity"]) != 0) for b in groups[code]]
        if old not in new:
            issue("现有配方投料差异", code, {"原配方": formula, "新BOM": groups[code], "新归属": owners[code]})
    ids = {r['name']: r['id'] for r in snapshot.get('formulas', [])}
    ids.update({code: ids.get(code, 'recipe:' + code) for code in groups})
    source_hashes = dict(db.execute("SELECT kind,sha256 FROM sources"))
    for code, blocks in groups.items():
        assigned = {(r['owner'], r['department']) for r in owners[code]}
        for block in blocks:
            active = [l for l in block['lines'] if not l['excluded']]
            samples = yields.get(code, [])
            yield_value = sum(v for n,v in samples) / len(samples) if samples else sum((Decimal(l['ratio']) for l in active), Decimal(0))
            formula = {'id': ids[code], 'name': code, 'kind': 'recipe', 'source_kind': 'erp_bom', 'source_row': block['row'], 'source_hashes': source_hashes,
                'sheet': 'Sheet', 'owner': next(iter(assigned))[0] if len(assigned)==1 else None,
                'department': next(iter(assigned))[1] if len(assigned)==1 else None,
                'yield': str(yield_value), 'yield_source': {'kind': 'mes_mean' if samples else 'bom_ratio_sum', 'count': len(samples), 'source_rows': [n for n,v in samples] if samples else [l['row'] for l in active]},
                'ratio_linked': False, 'lines': [dict(code=l['feed'], ref=ids[l['ref']] if l['ref'] in groups else l['ref'], kind='recipe' if l['ref'] in groups else 'material', quantity=l['quantity'], ratio=str(Decimal(l['ratio'])*100), source_row=l['row']) for l in active]}
            db.execute('INSERT INTO candidate_formulas VALUES(?,?,?)', (block['id'],code,packed(formula)))
    summary = {"created_at": datetime.now(UTC).isoformat(), "source_rows": {k: len(v) for k,v in rows.items()}, "bom_blocks": sum(map(len, groups.values())), "products": len(groups), "departments": dict(Counter(row[2] for n,row in rows["owners"])), "issues": dict(Counter(c for c,p,e in issues)), "inventory_pending": True}
    db.commit()
    assert db.execute("PRAGMA integrity_check").fetchone() == ("ok",)
    for kind, data in rows.items():
        assert db.execute("SELECT count(*) FROM raw_rows WHERE kind=?", (kind,)).fetchone()[0] == len(data)
        assert db.execute("SELECT source_row,values_json FROM raw_rows WHERE kind=? ORDER BY source_row", (kind,)).fetchall() == [(n, packed(r)) for n,r in data]
    db.close()
    (output / "分析摘要.json").write_text(packed(summary), encoding="utf-8")
    (output / "已确认规则.json").write_text(packed(RULES), encoding="utf-8")
    with (output / "待确认清单.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["序号", "类别", "编码", "证据", "确认结果"])
        for i,(category,code,evidence) in enumerate(issues,1):
            writer.writerow([i,category,code,packed(evidence),""])
    sections = []
    for category in summary["issues"]:
        entries = [(i,c,e) for i,(k,c,e) in enumerate(issues,1) if k == category]
        items = ''.join(f'<details><summary>{i}. {html.escape(c)}</summary>{render_evidence(e)}</details>' for i,c,e in entries)
        sections.append(f'<section><h2>{html.escape(category)}（{len(entries)}）</h2>{items}</section>')
    page = '''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>研发数据核对</title><style>body{font:16px/1.6 system-ui;margin:32px auto;max-width:1080px;padding:0 20px;color:#243740;background:#f5f7f7}section{background:white;padding:20px;margin:20px 0;border-radius:8px}h1{font-size:28px}h2{font-size:20px}summary{cursor:pointer;padding:10px;border-bottom:1px solid #ddd}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f5f7f7;padding:16px;font:14px/1.5 monospace}input{font:inherit;padding:10px;width:calc(100% - 24px)}details[hidden]{display:none}a{color:#174b6a}</style><h1>研发数据核对</h1><p>本页是候选核对结果。全部疑点确认、新库存复核及隔离验收完成后，才统一导入。原表与在线数据未修改。</p><p><a href="研发数据分析.db">分析数据库</a> · <a href="待确认清单.csv">完整待确认清单</a> · <a href="已确认规则.json">已确认规则</a></p><label for="search">搜索编码或证据</label><input id="search" placeholder="例如 B-52D、回收、研发五部"><main>CONTENT</main><script>document.getElementById('search').addEventListener('input',e=>{const q=e.target.value.toLowerCase();document.querySelectorAll('details').forEach(d=>d.hidden=!d.textContent.toLowerCase().includes(q))})</script></html>'''
    page = page.replace("</style>", "table{border-collapse:collapse;width:100%;font-size:14px}th,td{padding:8px;text-align:left;border-bottom:1px solid #ddd}th{background:#f5f7f7}h3{font-size:16px}</style>")
    (output / "研发数据核对.html").write_text(page.replace("CONTENT", ''.join(sections)), encoding="utf-8")
    return summary


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--snapshot", type=Path)
    args = parser.parse_args()
    print(packed(analyze(args.source, args.output, json.loads(args.snapshot.read_text(encoding="utf-8")) if args.snapshot else None)))
